from openpyxl import *

class ID_Team:

    import random
    
    def __init__(self,teamName):
        teamName += ".xlsx"
#        print teamName
        self.ws = []
        self.row_count = 1
        self.wb = load_workbook(filename = teamName, data_only=True)  #Load workbook
    
        self.ws = self.wb.active                              #switch to active worksheet 
    
        
    def Positions(self):
        from openpyxl import load_workbook
        
        self.row_count = self.ws.max_row
        
        for i in range(1,self.row_count):
            if self.ws.cell(row=i, column = 1).value == 'Runners':
                self.RunnersRow = i
        
            if self.ws.cell(row=i, column = 1).value == 'QBs':
                self.QBsRow = i
        
            if self.ws.cell(row=i, column = 1).value == 'Receivers':
                self.ReceiversRow = i
        
            if self.ws.cell(row=i, column=1).value == 'Kickers':
                self.KickersRow = i
            if self.ws.cell(row=i, column=1).value == 'Dline':
                self.DlineRow = i
            if self.ws.cell(row=i, column=1).value == 'LBs':
                self.LBsRow = i
            if self.ws.cell(row=i, column=1).value == 'DBs':
                self.DBsRow = i

        #Develop ranges and indices
        #Runners Range
        self.RunnersRange = range(self.RunnersRow+1,self.QBsRow-1)
        self.NumberOfRunners = self.QBsRow - 2 - self.RunnersRow
        self.RunnersStartIndex = self.RunnersRange[0]

        #QBs Range
        self.QBsRange = range(self.QBsRow+1,self.ReceiversRow-1)
        self.NumberOfQBs = self.ReceiversRow - 2  - self. QBsRow
        self.QBsStartIndex = self.QBsRange[0]

        #Receivers Range
        self.ReceiversRange = range(self.ReceiversRow+1,self.KickersRow-1)
        self.NumberOfReceivers = self.KickersRow - 2 - self.ReceiversRow
        self.ReceiversStartIndex = self.ReceiversRange[0]

#        print(self.ws["H2"].value)
        return self.ws
        





    
