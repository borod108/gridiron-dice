#-----------------------------------------------------------------------------
# Name:     RunYardage Table
# Purpose:  11x3 table indexed by roll of die. 
# Inputs:   rowIndex, columnIndex
# Outputs:  ["Loss",yards(int),"SG","MG","LG"]
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def RunYardageTable(rowIndex,columnIndex):
    if rowIndex == 0:
        yardageRow = ["SG","SG","LG"]
    if rowIndex == 1:
        yardageRow = ["Loss","Loss","SG"]
    if rowIndex == 2:
        yardageRow = [0,0,0]
    if rowIndex == 3:
        yardageRow = [1,2,3]
    if rowIndex == 4:
        yardageRow = [2,3,4]
    if rowIndex == 5:
        yardageRow = [3,4,5]
    if rowIndex == 6:
        yardageRow = [3,4,5]
    if rowIndex == 7:
        yardageRow = [5,6,7]
    if rowIndex == 8:
        yardageRow = [7,8,9]
    if rowIndex == 9:
        yardageRow = [8,9,"SG"]
    if rowIndex == 10:
        yardageRow = ["SG","SG","LG"]

    return yardageRow[columnIndex]

#-----------------------------------------------------------------------------
# Name:     Loss Table
# Purpose:  11x3 table indexed by roll of die. 
# Inputs:   rowIndex, columnIndex
# Outputs:  Loss yards ranging from -5 to -1
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def LossTable(rowIndex,columnIndex):
    if rowIndex == 0:
        Loss = [-5,-4,-3]
    if rowIndex == 1:
        Loss = [-4,-3,-2]
    if rowIndex == 2:
        Loss = [-3,-2,-1]
    if rowIndex == 3:
        Loss = [-2,-1,-1]
    if rowIndex == 4:
        Loss = [-1,-1,-1]
    if rowIndex == 5:
        Loss = [-1,-1,-1]
    if rowIndex == 6:
        Loss = [-1,-1,-1]
    if rowIndex == 7:
        Loss = [-2,-1,-1]
    if rowIndex == 8:
        Loss = [-3,-2,-1]
    if rowIndex == 9:
        Loss = [-4,-3,-2]
    if rowIndex == 10:
        Loss = [-5,-4,-3]
    return Loss[columnIndex]

#-----------------------------------------------------------------------------
# Name:     ShortGainTable
# Purpose:  11x3 table indexed by roll of die.  Provides short gain yardage 
# Inputs:   rowIndex, columnIndex - identify unique table value for yardage
# Outputs:  Short gains ranging from 10 to 20 yards
# Author:      Rick
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
#-----------------------------------------------------------------------------
def ShortGainTable(rowIndex,columnIndex):
    if rowIndex == 0:                       #Somewhat gaussian distribution
        yardageRow = [16,16,17]             #ranging from 10 to 20 - as an 
    if rowIndex == 1:                       #interesting exercise, will 
        yardageRow = [14,14,15]             #compute mean and standard        
    if rowIndex == 2:                       #deviation.
        yardageRow = [13,13,14]             #The way this table works is as
    if rowIndex == 3:                       #follows.  The calling routine
        yardageRow = [11,12,12]             #creates two uniformly-distributed
    if rowIndex == 4:                       #variables, a row index ranging 
        yardageRow = [10,10,11]             #from 0 to 10 and a column index
    if rowIndex == 5:                       #ranging from 0 to 2.  These 
        yardageRow = [10,10,10]             #indices are passed into this
    if rowIndex == 6:                       #11x3 table and the intersection
        yardageRow = [10,10,11]             #in the table that is indexed by
    if rowIndex == 7:                       #the two indices is extracted and
        yardageRow = [12,13,14]             #is the short gain yardage that is
    if rowIndex == 8:                       #returned to the calling routine
        yardageRow = [15,16,17]
    if rowIndex == 9:
        yardageRow = [17,18,18]
    if rowIndex == 10:
        yardageRow = [19,19,20]
    return yardageRow[columnIndex]  #Return the short gain yardage


#-----------------------------------------------------------------------------
# Name:     ShortPassTable
# Purpose:  11x3 table indexed by roll of die when the completed pass is a
#           short pass length
# Inputs:   rowIndex, columnIndex
# Outputs:  Short pass yards that results in anything from a 5 yard loss to a
#           long gain with the mean somewhere around 6 yards
# Author:      Rick
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def ShortPassTable(rowIndex,columnIndex):
    if rowIndex == 0:                       #This differs from the short gain
        yardageRow = ["SPG","SPG","LPG"]    #table in that there are some
    if rowIndex == 1:                       #string entries that require 
        yardageRow = ["Loss","Loss","MPG"]  #further resolution.
    if rowIndex == 2:
        yardageRow = [0,0,0]
    if rowIndex == 3:
        yardageRow = [1,2,3]
    if rowIndex == 4:
        yardageRow = [2,3,4]
    if rowIndex == 5:
        yardageRow = [3,4,5]
    if rowIndex == 6:
        yardageRow = [3,4,5]
    if rowIndex == 7:
        yardageRow = [5,6,7]
    if rowIndex == 8:
        yardageRow = [7,8,9]
    if rowIndex == 9:
        yardageRow = [8,9,"SPG"]
    if rowIndex == 10:
        yardageRow = ["SPG","MPG","LPG"]

    return yardageRow[columnIndex]




#-----------------------------------------------------------------------------
# Name:     MidPassTable
# Purpose:  11x3 table indexed by roll of die when the completed pass is of
#           medium length
# Inputs:   rowIndex, columnIndex
# Outputs:  Medium pass yards that results in anything from a 5 yard loss 
#           (probably need to fix this) to a long gain with the mean somewhere
#           around 12-13 yards.  This is the area that needs fixing so post
#           pattern receptions don't go for 3 yards.  Will have to have
#           a test so that short pass gains (SPG) don't result in a loss and
#           another test (possibly in a different method) that does not allow
#           long-medium pass completions to result in < 10 yards
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def MidPassTable(rowIndex,columnIndex):
    

    if rowIndex == 0:
        yardageRow = ["LPG","LPG","LPG"]    #LPG is long pass gain and uses
    if rowIndex == 1:                       #the LPGResolve() method
        yardageRow = ["LPG","SPG","LPG"]    #SPG is short pass gain and uses
    if rowIndex == 2:                       #the ShortPassGainTable
        yardageRow = [6,6,6]
    if rowIndex == 3:
        yardageRow = [8,8,8]
    if rowIndex == 4:
        yardageRow = [9,9,9]
    if rowIndex == 5:
        yardageRow = [10,10,10]
    if rowIndex == 6:
        yardageRow = [11,11,12]
    if rowIndex == 7:
        yardageRow = [12,13,13]
    if rowIndex == 8:
        yardageRow = [14,14,15]
    if rowIndex == 9:
        yardageRow = ["MPG","MPG","MPG"]    #MPG is the mid-length pass gain
    if rowIndex == 10:                      #and uses the MPGResolve() method
        yardageRow = ["LPG","LPG","LPG"]    #Same with LPG and LPGResolve()

    return yardageRow[columnIndex] #Return the yardage gained on the pass play



#-----------------------------------------------------------------------------
# Name:     MidLongPassTable
# Purpose:  11x3 table indexed by roll of die when the completed pass is from
#           a post pattern
# Author:      Rick
# Created:     3/27/2017
# Copyright:   (c) Rick 2017
# Parameters Passed In: rowIndex - Yardage Table row index
#                       columnIndex - Yardage Table column index
#-----------------------------------------------------------------------------
def MidLongPassTable(rowIndex,columnIndex):
    

    if rowIndex == 0:
        yardageRow = ["LPG","LPG","LPG"]    #Use LPGResolve()
    if rowIndex == 1:
        yardageRow = ["LPG","LPG","LPG"]
    if rowIndex == 2:
        yardageRow = [12,13,14]             #Most completed post patterns go 
    if rowIndex == 3:                       #for more than 10 yards
        yardageRow = [15,16,17]
    if rowIndex == 4:
        yardageRow = [18,19,20]
    if rowIndex == 5:
        yardageRow = [21,22,23]
    if rowIndex == 6:
        yardageRow = [24,25,26]
    if rowIndex == 7:
        yardageRow = [27,28,29]
    if rowIndex == 8:
        yardageRow = [30,31,32]
    if rowIndex == 9:
        yardageRow = [33,34,35]
    if rowIndex == 10:
        yardageRow = ["LPG","LPG","LPG"]
    

    return yardageRow[columnIndex]  #Return yardage gained on the pass play




#-----------------------------------------------------------------------------
# Name:     ShortPassGainTable
# Purpose:  11x3 table indexed by roll of die when the yardage result is an
#           "SPG"
# Inputs:   rowIndex, columnIndex
# Outputs:  Yardage result ranges from 16 yards to a long gain with a mean 
#           somewhere around 20 to 21 yards
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def ShortPassGainTable(rowIndex,columnIndex):
    
    if rowIndex == 0:
        yardageRow = ["MPG","MPG","MPG"]    #Use MPGResolve()
    if rowIndex == 1:
        yardageRow = [21,21,20]
    if rowIndex == 2:
        yardageRow = [20,20,19]
    if rowIndex == 3:
        yardageRow = [18,18,17]
    if rowIndex == 4:
        yardageRow = [16,16,17]
    if rowIndex == 5:
        yardageRow = [17,18,18]
    if rowIndex == 6:
        yardageRow = [19,19,20]
    if rowIndex == 7:
        yardageRow = [20,21,21]
    if rowIndex == 8:
        yardageRow = [22,22,23]
    if rowIndex == 9:
        yardageRow = [23,23,24]
    if rowIndex == 10:
        yardageRow = ["MPG","MPG","MPG"]    #MPGResolve()

    return yardageRow[columnIndex]  #Return yardage gained on pass play



#-----------------------------------------------------------------------------
# Name:     SackYardageTable
# Purpose:  11x3 table indexed by roll of die after a sack has occurred 
# Inputs:   rowIndex, columnIndex
# Outputs:  After a sack occurs, used to determine the extent of the sack
#           ranging from -16 to 0 yards with the mean somewhere around - 6 or
#           -7 yards
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def SackYardageTable(rowIndex,columnIndex):
    if rowIndex == 0:
        yardageRow = [-16,-15,-15]
    if rowIndex == 1:
        yardageRow = [-14,-14,-13]
    if rowIndex == 2:
        yardageRow = [-13,-12,-12]
    if rowIndex == 3:
        yardageRow = [-11,-11,-10]
    if rowIndex == 4:
        yardageRow = [-10,-9,-9]
    if rowIndex == 5:
        yardageRow = [-8,-8,-7]
    if rowIndex == 6:
        yardageRow = [-7,-6,-6]
    if rowIndex == 7:
        yardageRow = [-5,-5,-4]   
    if rowIndex == 8:
        yardageRow = [-4,-3,-3]
    if rowIndex == 9:
        yardageRow = [-2,-2,-1]
    if rowIndex == 10:
        yardageRow = [-1,0,0]

    return yardageRow[columnIndex]  #Return sack yardage



#-----------------------------------------------------------------------------
# Name:     ResultofthePass
# Purpose:  Called from ResultofthePlay, determines whether the pass results
#           in a sack, completion, incompletion or an INT
# Inputs:   PassLength: Short, medium or Long
#           PC: QB's percent completion
#           PI: QB's percent interceptions
#           SackPerc: The percentage of sacks allowed by the offense
#           Dstats: A 3-element list consisting of average rush yards allowed 
#           by the defense, percentage completions allowed by the defense, and
#           yards after catch allowed by the defense
# Outputs:  Sack, completion, incompletion or an INT
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def ResultofthePass(PassLength,PC,PI,SackPerc,Dstats):
    import random


#Based upon middle 10 teams from 2016 NCAA statistics   
    AvePassCompPerc = 59.1
    PassCompletedRatio = (PC/AvePassCompPerc)*(Dstats[1]/AvePassCompPerc)
    
    
#Pass completion average depends on the QBs individual stats and the length of
#the pass.  Same with the INT percentage
    if PassLength == "Short":
        PassCompPerc = int(78 * PassCompletedRatio)
        
        PI = PI - 1
    if PassLength == "Mid":
        PassCompPerc = int(67 * PassCompletedRatio) #Bumped up average pass
    if PassLength == "MidLong":                     #completion % by 2 to 
        PassCompPerc = int(49 * PassCompletedRatio) #account for the new pass
                                                    #MidLong
    if PassLength == "Long":
        PassCompPerc = int(35 * PassCompletedRatio)
        
        PI = PI + 1
      

    #Test to see if a sack occurred.
    dice1 = random.randint(0,9999)
    if dice1 <= 100*SackPerc:
        SackYards = 7
        return "Sack"
    else:

        dice1 = random.randint(0,99)
        if dice1 <= PI:                 #See if pass is intercepted
            return "Int"
        if (dice1 > PI) and (dice1 <= (PassCompPerc+PI)):
            return "Completed"
        else:
            return "Incomplete"


#-----------------------------------------------------------------------------
# Name:     ResolveGains
# Purpose:  If the yardageResult is a loss, SG, MG, LG, SPG or MPG, this is
#           where the codes get turned into actual yards, often using the 
#           yardage tables above
# Inputs:   yardageResult: if not a numerical value, coded as SG, MG, LG, SPG 
#           or MPG
# Outputs:  Numerical value of yards
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def ResolveGains(yardageResult):
    import random
    import Die

    #This is where Loss, SG, MG, LG gets resolved
    if yardageResult == "Loss":     #Use the above LossTable to resolve losses
        lossRoll=Die.RolltheDie()
        all3die = lossRoll.Dice3Dice6Dice6()
        dice1,dice2,dice3 = lossRoll.dice1,lossRoll.dice2,lossRoll.dice3
        rowIndex = dice2 + dice3 - 2
        columnIndex = dice1 - 1
        yardageResult = LossTable(rowIndex,columnIndex)
        
    if yardageResult == "SG":        #Use the above ShortGainTable to resolve   
        SGRoll=Die.RolltheDie()             #short gains
        all3die = SGRoll.Dice3Dice6Dice6()
        dice1,dice2,dice3 = SGRoll.dice1,SGRoll.dice2,SGRoll.dice3
        rowIndex = dice2 + dice3 -2
        columnIndex = dice1 - 1       
        yardageResult = ShortGainTable(rowIndex,columnIndex)    #Access the short gain table
        
    if yardageResult == "MG":       #MG is resolved using 3 dice to produce yardage results 
        MGRoll = Die.RolltheDie()                           #ranging from 17 to 52 yards    
        yardageResult = MGRoll.SingleValue366()  #Access the medium gain table
    if yardageResult == "LG":                              #LG is resolved using 3 dice to produce
        LGRoll = Die.RolltheDie()                           #yardage results ranging from 17 to 99 yards

        yardageResult = LGRoll.SingleValue31010()   #Access the long gain table (3 dice, 
                                                                             # 10 sides to a die)

    if yardageResult == "SPG":      #SPG is resolved using the 
        SPGRoll=Die.RolltheDie()    #ShortPassGainTable

#roll 1 3-sided die and 2 6-sided dice        
        all3die = SPGRoll.Dice3Dice6Dice6()
        dice1,dice2,dice3 = SPGRoll.dice1,SPGRoll.dice2,SPGRoll.dice3
        
        rowIndex = dice2 + dice3 - 2    #Row and column indices for the short 
        columnIndex = dice1 - 1         #pass yardage table.

#Index short pass gain yardage table to extract yardage result
        yardageResult = ShortPassGainTable(rowIndex,columnIndex)

    if yardageResult == "MPG":          #MPG is resolved using MPGResolve()
        yardageResult = MPGResolve()    #Medium pass gain yardage is resolved
                                        #using a linear combination of three
                                        #die
    if yardageResult == "LPG":          #LPG is resolved using LPGResolve()
        yardageResult = LPGResolve()    #Same linear combination as 
                                        #MPGResolve()
        
    return yardageResult
    


#-----------------------------------------------------------------------------
# Name:     MPGResolve
# Purpose:  If the yardageResult is an MPG, this resolves the yardage by using
#           3 dice to yield values ranging from 2 to 54 yards.
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def MPGResolve():
    import random
    import Die
    
    MPGRoll = Die.RolltheDie()
    return MPGRoll.SingleValue577()


#-----------------------------------------------------------------------------
# Name:     LPGResolve
# Purpose:  If the yardageResult is an LPG, this resolves the yardage by using
#           3 dice to yield values ranging from 17 to 99 yards.
# Inputs:   None
# Outputs:  Numerical value of yards
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def LPGResolve():
    import random
    import Die
    
    LPGRoll = Die.RolltheDie()
    yards = LPGRoll.SingleValue91010()
    
    return yards

#-----------------------------------------------------------------------------
# Name:     KickReturn
# Purpose:  If a kickoff is returned, this method computes the return yardage
#           based upon the returner's average and longest gain
# Inputs:   TeamReceivingtheKick
# Outputs:  Kick return in yards
# Author:      Rick
#
# Created:     10/20/2016
# Copyright:   (c) Rick 2016
# Licence:     <your licence>
#-----------------------------------------------------------------------------
def KickReturn(TeamReceivingtheKick):
    import openpyxl
    import PickAPlayer
    import random
    
    global KickReturner #Name of the player returning the kick
    
#    KickReturnerRow = 52    #Where the kick returner stats are stored
    KickReturnSigma = 5     #Estimated standard deviation on kick returns
    KickReturnLGAdder = 30  #Take the kick returners longest gain and add 30
    KickReturnerColumn = 2  #yards to it to set the long gain maximum
    RequestedKickReturnerRow = PickAPlayer.PickAPlayer(TeamReceivingtheKick,0,
                                                       0,0,0,0) 
    RequestedKickReturnerRow.FindStats()
    KickReturnerRow = RequestedKickReturnerRow.KRsPosition + \
        RequestedKickReturnerRow.Offset
 
#Extract kick returner stats    
    KickReturnAve = round(TeamReceivingtheKick.cell(row = KickReturnerRow,
                                             column = 4).value,0)
    KickReturnLG = TeamReceivingtheKick.cell(row = KickReturnerRow,
                                             column = 5).value
    KickReturner = TeamReceivingtheKick.cell(row = KickReturnerRow,
                                            column = KickReturnerColumn).value
    
    FirstReturnTest = random.randint(1,10)  #~10% of returns are short
    KickReturn = 25                         #Initialization value.
    if FirstReturnTest == 1:                #Short return ranging between 5 
        KickReturn = random.randint(5,15)   #and 15 yards
        
#80% of returns are the returners average +/- 1 sigma value
    elif FirstReturnTest <= 9:
        KickReturn = random.randint(KickReturnAve - KickReturnSigma, 
                                    KickReturnAve + KickReturnSigma)

#Remaining 10% of returns go from returners average + 1 sigma to the
#returner's longest return + the long gain adder constant defined above
    if FirstReturnTest == 10:
        KickReturn = random.randint(KickReturnAve + KickReturnSigma,
                                    KickReturnLG + KickReturnLGAdder)
    
    return KickReturn


#-----------------------------------------------------------------------------
# Name:     YardlineAdjust
# Purpose:  When the yardline crosses the 50 yardline, the yardline decreases
# Inputs:   Yardline (1 - 100)
# Outputs:  Adjusted Yardline (1 - 50 - 0 with <=0 being a touchdown
# Author:      Rick
#
# Created:     1/5/2017
# Copyright:   (c) Rick 2017
#-----------------------------------------------------------------------------
def YardlineAdjust(Yardline):
    
    AdjustedYardline = 100 - Yardline
    return AdjustedYardline


#-----------------------------------------------------------------------------
# Name:     Touchback Test
# Purpose:  When the YardLine > 100 after a punt, it's a touchback
# Inputs:   Yardline (1 - 100)
# Outputs:  Flag saying it's a touchback
# Author:      Rick
#
# Created:     1/9/2017
# Copyright:   (c) Rick 2017
#-----------------------------------------------------------------------------
def TouchbackTest(Yardline):
    
    TouchbackFlag = 0
    if Yardline > 99:
        TouchbackFlag = 1
    return TouchbackFlag



#-------------------------------------------------------------------------------
# Name:     PuntReturn
# Purpose:  If a punt is returned, this method computes the return yardage
#           based upon the returner's average and longest gain
# Inputs:   TeamReceivingtheKick (really the punt)
# Outputs:  Punt return in yards
# Author:      Rick Burney
#
# Created:     2/24/2017
# Copyright:   (c) Rick 2017
# Licence:     <your licence>
#-------------------------------------------------------------------------------
def PuntReturn(TeamReceivingtheKick,Force,PunterFCPerc):
    import openpyxl
    import PickAPlayer
    import random
    
    global FCFlag
    global PuntReturner

    
    FCFlag = 0  #Initialize the Fair Catch Flag
   
    PuntReturnerRow = 55      #Where the punt returner stats are located
    PuntReturnSigma = 8        #1-sigma on average return
    PuntReturnLGAdder = 10  #Long gain adder.  
    PuntReturnerColumn = 2
    PuntReturnFCColumn = 3
    PuntReturnAveColumn = 4
    PuntReturnAveLGColumn = 5
    
    RequestedPuntReturnerRow = PickAPlayer.PickAPlayer(TeamReceivingtheKick,0,
                                                       0,0,0,0) 
    RequestedPuntReturnerRow.FindStats()
    
#Use kick returners row location but increase the offset to get to the PR row
    PuntReturnerRow = RequestedPuntReturnerRow.KRsPosition + \
        RequestedPuntReturnerRow.PROffset
    
#On 12/19/18, made the change to have fair catch percentage determined by the
#Punter.  As a consequence, we are locating the punter stats which is on the 
#kicking team worksheet inspite of the fact that this is a punt return method
#Extract punt returner stats
    PuntReturnFCPerc = PunterFCPerc #This comes from the punter

    PuntReturnAve = round(TeamReceivingtheKick.cell(row = PuntReturnerRow,
                                        column = PuntReturnAveColumn).value,0)
    PuntReturnLG = TeamReceivingtheKick.cell(row = PuntReturnerRow,
                                        column = PuntReturnAveLGColumn).value
    PuntReturner = TeamReceivingtheKick.cell(row = PuntReturnerRow,
                                            column = PuntReturnerColumn).value

    FCTest = random.randint(0,99)   #Test for Fair Catch.  Fair catch # from
    if FCTest <= PuntReturnFCPerc:  #stats sheet was pulled somewhere from the
        PuntReturn = 0              #web, can't tell you where        
        FCFlag = 1      #Set fair catch flag
       
#The following assumptions about punt returns I pulled out of my ass
    else:                                       #If not fair catch    
        FirstReturnTest = random.randint(1,10)  #~20% of punt returns are not
        PuntReturn = 0                          #returned
        
        if FirstReturnTest <= 2:    #This is the 20% part
            PuntReturn = 0
        elif FirstReturnTest == 3:              #10% of returns range from 0 
            PuntReturn = random.randint(0,15)   #to 15 yards, independent of
                                                #punt returner's stats

#60% of the returns are based upon the punt returner's average, +/- 1-sigma
        elif FirstReturnTest <= 9:
            PuntReturn = random.randint(PuntReturnAve - PuntReturnSigma, 
                                        PuntReturnAve + PuntReturnSigma)

#10% of returns are from the average + 1-sigma to the returner's long gain +
#long gain adder.  The long gain adder is intended to allow a punt returner to
#return one for a TD, even if that did not happen during the stat year
        elif FirstReturnTest <= 10:
            PuntReturn = random.randint(PuntReturnAve + PuntReturnSigma,
                                            PuntReturnLG + PuntReturnLGAdder)
            
            
#Return a 2-element list, 1st element is the return yardage, 2nd element is 
#the name of the guy returning the punts
    PuntReturnList = [PuntReturn,PuntReturner]    
    return PuntReturnList