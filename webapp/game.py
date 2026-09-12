"""Headless game controller.

This is ControlPanel.py (Rev 4.0) with the Tkinter widgets removed.  Every
button handler of the desktop app becomes a method here; the module-level
globals of the original become attributes on the Game instance.  The engine
modules under ./engine are the original simulator files, untouched except for
Python 3 compatibility, and they talk to the UI through engine/bus.py.

All Excel I/O in the engine uses paths relative to the current working
directory, exactly like the desktop app, so the web server chdir()s into the
data directory once at startup.
"""
import glob
import json
import os
import secrets
import sys
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.join(HERE, "engine")
if ENGINE not in sys.path:
    sys.path.insert(0, ENGINE)

from openpyxl import Workbook, load_workbook   # noqa: E402
import bus                                      # noqa: E402
import GeneralInfoFile                          # noqa: E402
import LogPlays                                 # noqa: E402
import PickAPlayer                              # noqa: E402
import Play                                     # noqa: E402
import random                                   # noqa: E402
import history                                  # noqa: E402

OUTPUT_SUFFIXES = ("Log.xlsx", "log.xlsx", "Stats.xlsx", "DStats.xlsx")

# Checkbox names exactly as on the desktop scoreboard
CHECKBOXES = [
    ("RunOriented", "Run-Oriented"), ("PassOriented", "Pass-Oriented"),
    ("HurryUp", "Hurry-Up"), ("Spike", "Spike the Ball"),
    ("HailMaryPass", "Hail Mary Pass"), ("QBTakesaKnee", "Take a Knee"),
    ("Squib", "Squib Kick"), ("Onside", "Onside Kick"), ("FK", "Free Kick"),
    ("PlacementPunt", "Placement Punt"),
    ("HTBlowout", "HT Blowout"), ("VTBlowout", "VT Blowout"),
    ("ForcePlay", "Force (debug)"),
]
KICK_BOXES = ("Squib", "Onside", "FK", "PlacementPunt", "RunOriented",
              "PassOriented", "HurryUp", "Spike", "HailMaryPass")


def list_team_files(data_dir):
    """Team workbooks = .xlsx files that are not simulator outputs."""
    teams = []
    for f in sorted(os.listdir(data_dir)):
        if not f.lower().endswith(".xlsx") or f.startswith("~$"):
            continue
        if f.endswith(OUTPUT_SUFFIXES) or f == "Team Selection.xlsx":
            continue
        teams.append(f[:-5])
    return teams


class GameError(Exception):
    pass


class Game:
    """One game in progress.  Construct = press Start."""

    def __init__(self, home, visitor, seed=None, gid=None):
        if home == visitor:
            raise GameError("Home and visiting team must differ")
        for t in (home, visitor):
            if not os.path.exists(t + ".xlsx"):
                raise GameError("Missing team file %s.xlsx" % t)
        self.homeTeamName = home
        self.visitingTeamName = visitor
        self.boxes = {k: 0 for k, _ in CHECKBOXES}
        self.log = []            # play-by-play lines shown in the UI
        self.finished = False
        self.output_files = []
        self.last_turn = None
        self.last_quit = None
        self.messages = []
        self.id = gid or history.new_id(home, visitor)
        history.start(self.id, home, visitor)
        # Every random number the engine draws comes from the global `random`
        # module, so seeding it here and journaling each button press makes the
        # whole game reproducible: see replay(), undo() and resume_latest().
        self.seed = seed if seed is not None else secrets.randbits(63)
        self.journal = []
        self._replaying = False
        random.seed(self.seed)
        bus.reset_all()
        self._start_game()
        self.save_journal("in_progress")

    # ---------------------------------------------------------------- journal
    JOURNAL = "journal.json"

    def journal_path(self):
        return os.path.join(history.game_dir(self.id), self.JOURNAL)

    def save_journal(self, status):
        if self._replaying:
            return
        os.makedirs(history.game_dir(self.id), exist_ok=True)
        tmp = self.journal_path() + ".tmp"
        with open(tmp, "w") as f:
            json.dump({"id": self.id, "home": self.homeTeamName, "visitor": self.visitingTeamName,
                       "seed": self.seed, "status": status, "actions": self.journal}, f)
        os.replace(tmp, self.journal_path())

    @classmethod
    def replay(cls, j, upto=None):
        """Rebuild a game from a journal dict, applying its first `upto` actions."""
        g = cls(j["home"], j["visitor"], seed=j["seed"], gid=j["id"])
        g._replaying = True
        try:
            for action, boxes in j["actions"][:upto]:
                g.last_turn = g.do(action, {k: 1 for k in boxes})
        finally:
            g._replaying = False
        g.save_journal("in_progress")
        return g

    def undo(self):
        """Return a new Game equal to this one minus its last action."""
        if not self.journal:
            raise GameError("Nothing to undo")
        j = {"id": self.id, "home": self.homeTeamName, "visitor": self.visitingTeamName,
             "seed": self.seed, "actions": self.journal[:-1]}
        return Game.replay(j)

    @classmethod
    def resume_latest(cls):
        """After a restart: rebuild the most recent in-progress game, if any."""
        candidates = []
        for path in glob.glob(os.path.join(history.GAMES_DIR, "*", cls.JOURNAL)):
            try:
                with open(path) as f:
                    j = json.load(f)
            except (OSError, ValueError):
                continue
            if j.get("status") == "in_progress":
                candidates.append(j)
        if not candidates:
            return None
        j = max(candidates, key=lambda j: j["id"])
        return cls.replay(j)

    # ------------------------------------------------------------------ setup
    def _load(self, name):
        return load_workbook(filename=name + ".xlsx", data_only=True).active

    def _ensure_log_file(self, name):
        fn = name + "Log.xlsx"
        if not os.path.exists(fn):
            wb = Workbook()
            wb.active.cell(row=1, column=1).value = "Play Type"
            wb.save(fn)

    def _positions(self, ws):
        """ControlPanel.Positions(): QB names + section start rows."""
        rr = PickAPlayer.PickAPlayer(ws, 0, 0, 0, 0, 0)
        rr.FindStats()
        qbs = rr.QBsPosition + rr.Offset
        return {
            "RunnersStartIndex": rr.RunnersPosition + rr.Offset,
            "QBsStartIndex": qbs,
            "ReceiversStartIndex": rr.ReceiversPosition + rr.Offset,
            "QBName": ws.cell(row=qbs, column=2).value,
            "BackupQBName": ws.cell(row=qbs + 1, column=2).value,
        }

    def _start_game(self):
        DRushTeamAveColumn, DSacksPerColumn = 3, 6
        DIntPercColumn, DFumRecoveryColumn = 10, 10
        PenaltyStatsColumn, ConferenceFactorColumn = 2, 3

        self.visitingTeam = self._load(self.visitingTeamName)
        self.homeTeam = self._load(self.homeTeamName)
        self.HomeTeamLogName = self.homeTeamName + "log.xlsx"
        self.VisitingTeamLogName = self.visitingTeamName + "log.xlsx"

        self._ensure_log_file(self.homeTeamName)
        self._ensure_log_file(self.visitingTeamName)
        LogPlays.ClearLog(self.homeTeamName, self.visitingTeamName)

        self.HTMDA, self.VTMDA, self.HTDMDA, self.VTDMDA = [], [], [], []

        H = PickAPlayer.PickAPlayer(self.homeTeam, 0, 0, 0, 0, 0)
        H.FindStats()
        V = PickAPlayer.PickAPlayer(self.visitingTeam, 0, 0, 0, 0, 0)
        V.FindStats()

        def dstats(ws, rr):
            base = rr.DStatsPosition
            return [ws.cell(row=base + rr.Offset, column=DRushTeamAveColumn).value,
                    ws.cell(row=base + rr.DPCompPercOffset, column=DRushTeamAveColumn).value,
                    ws.cell(row=base + rr.DPYACOffset, column=DRushTeamAveColumn).value,
                    ws.cell(row=base + rr.DPCompPercOffset, column=DSacksPerColumn).value,
                    ws.cell(row=base + rr.Offset, column=DIntPercColumn).value,
                    ws.cell(row=base + rr.DPCompPercOffset, column=DFumRecoveryColumn).value]

        # Same naming inversion as the original: VDstats holds the HOME sheet's
        # defensive numbers (used when the visitor is on offense) and vice versa.
        self.VDstats = dstats(self.homeTeam, H)
        self.HDstats = dstats(self.visitingTeam, V)

        vpen = V.TeamStatsPosition + V.Offset
        hpen = H.TeamStatsPosition + H.Offset
        self.VTeamStats = [self.visitingTeam.cell(row=vpen, column=PenaltyStatsColumn).value]
        self.HTeamStats = [self.homeTeam.cell(row=hpen, column=PenaltyStatsColumn).value]
        self.HomeTeamConferenceFactor = self.homeTeam.cell(row=hpen, column=ConferenceFactorColumn).value
        self.VisitingTeamConferenceFactor = self.visitingTeam.cell(row=vpen, column=ConferenceFactorColumn).value

        vp = self._positions(self.visitingTeam)
        hp = self._positions(self.homeTeam)
        self.VQBName, self.VBackupQBName = vp["QBName"], vp["BackupQBName"]
        self.HQBName, self.HBackupQBName = hp["QBName"], hp["BackupQBName"]
        # The desktop app called Positions() for the visitor then the home team,
        # so the globals ended up pointing at the HOME sheet's rows.
        self.RunnersStartIndex = hp["RunnersStartIndex"]
        self.QBsStartIndex = hp["QBsStartIndex"]
        self.ReceiversStartIndex = hp["ReceiversStartIndex"]

        self._coin_toss()
        bus.board["clock"] = "15:00"
        GeneralInfoFile.WriteGeneralInfo(self.homeTeamName, self.visitingTeamName)

        self.GM = {'Down': 1, 'YTG': 10, 'YardLine': 25, 'AdjustedYardLine': 25, 'TimeLeftinQuarter': 900, 'Quarter': 1,
                   'Offense': self.TeamWiththeBall, 'Defense': self.TeamthatDoesNotHavetheBall, 'OffenseFlag': self.TeamWiththeBallFlag,
                   'OldYardLine': 25, 'OldAdjustedYardline': 25, 'OldDown': 1, 'OldYTG': 10, 'OldTimeLeftInQuarter': 900,
                   'OldQuarter': 1, 'HomeTeamScore': 0, 'VisitingTeamScore': 0, 'CoPFlag': 0, 'OldHTS': 0, 'OldVTS': 0,
                   'Touchback': 0, 'ConversionOnDowns': 0, 'TimeCode': 0, 'OTFlag': 0, 'OTCounter': 0, 'SafetyFlag': 0,
                   'DTDFlag': 0, 'homeTeamPlayCount': 0, 'visitingTeamPlayCount': 0, 'OTSeries': 1, 'OTPossession': 1,
                   'TDFlag': 0, 'UntimedDownFlag': 0, 'HomeTimeouts': 3, 'VisitorTimeouts': 3, 'OldTimeLeftInQuarter': 0,
                   'TurnoverTD': 0, 'ClockRunning': 0, 'ConversionFlag': 0}
        self.GameOptions = {'RunCentric': 0, 'PassCentric': 0, 'Hup': 0, 'SpiketheBall': 0, 'HailMary': 0, 'Blowout': 0,
                            'ForcePenalty': 0, 'QBTakesaKnee': 0, 'ForcePlay': 0}
        self.Kicking = {'KickFlag': 0, 'StartoftheGame': 1, 'Startofthe3rdQuarter': 0, 'OnsideKick': 0, 'SquibKick': 0,
                        'FreeKick': 0, 'ThirdQuarterKickoffTeam': self.TeamthatDoesNotHavetheBall,
                        'ThirdQuarterTeamReceivingtheKick': self.TeamWiththeBall, 'ThirdQuarterOffenseFlag': 0, 'XPtFlag': 0,
                        'KickoffFlag': 0, 'FGFlag': 0, 'PlacementPunt': 0, 'PuntFlag': 0, 'TwoPointConvFlag': 0,
                        'ThirdQuarterKickoffRow': 0}
        self.PlayerStats = {'RunnerSI': self.RunnersStartIndex, 'ReceiverSI': self.ReceiversStartIndex,
                            'QBSI': self.QBsStartIndex}
        self.TeamStats = self._team_stats()
        self.log.append(bus.board.get("last_result", ""))
        self.messages = []

    def _coin_toss(self):
        if random.randint(1, 2) == 1:
            bus.results.append("Home Team will receive the kickoff")
            bus.board["last_result"] = "Home Team will receive the kickoff"
            self.TeamWiththeBall, self.TeamWiththeBallFlag = self.visitingTeam, 1
            self.TeamthatDoesNotHavetheBall = self.homeTeam
        else:
            bus.results.append("Visiting Team will receive the kickoff")
            bus.board["last_result"] = "Visiting Team will receive the kickoff"
            self.TeamWiththeBall, self.TeamWiththeBallFlag = self.homeTeam, 0
            self.TeamthatDoesNotHavetheBall = self.visitingTeam

    def _team_stats(self):
        return {'HomeDStats': self.HDstats, 'VisitorDStats': self.VDstats,
                'HomePenaltiesPerGame': self.HTeamStats[0],
                'VisitorPenaltiesPerGame': self.VTeamStats[0],
                'HTCF': self.HomeTeamConferenceFactor,
                'VTCF': self.VisitingTeamConferenceFactor,
                'homeTeamName': self.homeTeamName,
                'visitingTeamName': self.visitingTeamName,
                'PreCoPOffenseFlag': self.GM['OffenseFlag'] if hasattr(self, "GM") else self.TeamWiththeBallFlag,
                'HomeTeamBlowout': self.boxes["HTBlowout"],
                'VisitingTeamBlowout': self.boxes["VTBlowout"],
                'homeTeam': self.homeTeam, 'visitingTeam': self.visitingTeam}

    def _play(self):
        return Play.Play(self.GM, self.PlayerStats, self.GameOptions, self.TeamStats, None, None, self.Kicking)

    def _take_logs(self, p):
        self.HTMDA, self.VTMDA, self.HTDMDA, self.VTDMDA = p.HTMDA, p.VTMDA, p.HTDMDA, p.VTDMDA

    def _reset_boxes(self):
        for k in KICK_BOXES:
            self.boxes[k] = 0

    # ---------------------------------------------------------------- actions
    ACTIONS = ("call_play", "kickoff", "fg", "xpt", "go_for_two", "punt",
               "hto", "vto", "ot")

    def do(self, action, boxes):
        """Run one button press.  Returns dict(messages, results, error)."""
        if self.finished:
            raise GameError("Game is over")
        if action not in self.ACTIONS:
            raise GameError("Unknown action %s" % action)
        for k in self.boxes:
            self.boxes[k] = 1 if boxes.get(k) else 0
        self.journal.append([action, sorted(k for k in self.boxes if self.boxes[k])])
        self.save_journal("in_progress")
        bus.reset_turn()
        error = None
        try:
            getattr(self, action)()
        except Exception:
            error = traceback.format_exc()
        for r in bus.results:
            self.log.append(r)
        self.messages = list(bus.messages)
        return {"messages": self.messages, "results": list(bus.results), "error": error}

    def call_play(self):
        GM, GO, K = self.GM, self.GameOptions, self.Kicking
        if ((GM['Quarter'] == 1) or (GM['Quarter'] == 3)) and (GM['TimeLeftinQuarter'] == 900):
            bus.messages.append(("Button Press Infraction", "Press the Kickoff Button"))
            return
        if GM['ConversionFlag'] == 1:
            bus.messages.append(("Button Press Infraction",
                                 "Press Extra Point, Go For 2 button or Kickoff button as appropriate"))
            return
        if K['KickoffFlag'] == 1:
            bus.messages.append(("Button Press Infraction", "Only press the Kickoff Button"))
            return
        GO['Blowout'] = 0
        self.PlayerStats = {'RunnerSI': self.RunnersStartIndex, 'ReceiverSI': self.ReceiversStartIndex,
                            'QBSI': self.QBsStartIndex}
        b = self.boxes
        GO['RunCentric'] = b["RunOriented"]
        GO['PassCentric'] = b["PassOriented"]
        GO['Hup'] = b["HurryUp"]
        GO['SpiketheBall'] = b["Spike"]
        GO['HailMary'] = b["HailMaryPass"]
        GO['ForcePlay'] = b["ForcePlay"]
        GO['QBTakesaKnee'] = b["QBTakesaKnee"]
        if GO['RunCentric'] == 1 and GO['PassCentric'] == 1:
            b["PassOriented"] = b["RunOriented"] = 0
            GO['RunCentric'] = GO['PassCentric'] = 0
        if GM['OffenseFlag'] == 0 and b["HTBlowout"] == 1:
            GO['Blowout'] = 1
        if GM['OffenseFlag'] == 1 and b["VTBlowout"] == 1:
            GO['Blowout'] = 1
        self.TeamStats = self._team_stats()
        self.PlayerStats['QBSI'] = self.QBsStartIndex + (1 if GO['Blowout'] == 1 else 0)
        GO['ForcePenalty'] = GO['ForcePlay']

        p = self._play()
        p.OnOffenseIndication()
        p.PlayCall()
        p.PreSnapPenaltyTest()
        PenaltyList = p.PenaltyList
        if PenaltyList[1] == 0:
            p.PlayResult()
            ResultList = p.ResultList
            if PenaltyList[0] in ("Pass Interference on the Defense",
                                  "Pass Interference on the Offense") and ResultList[2] == "Sack":
                PenaltyList[5] = 0
            if p.FumbleResult[0] == 0 and p.IntResult[5] == 0:
                p.UpdateYardLine()
            if PenaltyList[5] == 1:
                p.PenaltyProcessing()
            p.GameManagement()
            self.GM = p.GM
            if self.GM['CoPFlag'] == 1 and self.GM['OTFlag'] == 0:
                p.CoP()
                self._reset_boxes()
        p.DisplayManagement()
        p.LogPlay(self.HTMDA, self.VTMDA, self.HTDMDA, self.VTDMDA)
        self._take_logs(p)
        self.GM['SafetyFlag'] = 0
        self.GM['ScoreFlag'] = 0
        p.TGraphics()

    def kickoff(self):
        GM, GO, K = self.GM, self.GameOptions, self.Kicking
        GM['ConversionFlag'] = 0
        self.TeamStats = self._team_stats()
        GO['ForcePlay'] = self.boxes["ForcePlay"]
        K['KickoffFlag'] = 1
        p = self._play()
        K['SquibKick'] = self.boxes["Squib"]
        K['OnsideKick'] = self.boxes["Onside"]
        K['FreeKick'] = self.boxes["FK"]
        if K['SquibKick'] + K['OnsideKick'] + K['FreeKick'] > 1:
            self.boxes["Squib"] = self.boxes["Onside"] = self.boxes["FK"] = 0
            K['SquibKick'] = K['OnsideKick'] = K['FreeKick'] = 0
            bus.messages.append(("Dummy", "Only one of the kicking checkboxes can be checked at once so all were reset and ignored"))
        p.Kickoff()
        p.LogPlay(self.HTMDA, self.VTMDA, self.HTDMDA, self.VTDMDA)
        self._take_logs(p)
        p.GameManagement()
        p.DisplayManagement()
        p.TGraphics()
        GM['ConversionFlag'] = 0
        self._reset_boxes()

    def fg(self):
        GM, GO, K = self.GM, self.GameOptions, self.Kicking
        GO['ForcePlay'] = self.boxes["ForcePlay"]
        if GM['ConversionFlag'] == 1:
            bus.messages.append(("Button Press Infraction",
                                 "Press Extra Point, Go For 2 button or Kickoff button as appropriate"))
            return
        p = self._play()
        p.FieldGoal()
        if K['FGFlag'] == 1:
            p.GameManagement()
            p.DisplayManagement()
            p.TGraphics()
            p.LogPlay(self.HTMDA, self.VTMDA, self.HTDMDA, self.VTDMDA)
            self._take_logs(p)
            K['FGFlag'] = 0
        K['KickFlag'] = 0

    def xpt(self):
        GM, GO, K = self.GM, self.GameOptions, self.Kicking
        if GM['OTFlag'] == 1 and GM['OTSeries'] >= 3:
            bus.messages.append(("Overtime", "On or after the 3rd series, and after a TD, the scoring team must go for 2 points"))
            return
        GO['ForcePlay'] = self.boxes["ForcePlay"]
        p = self._play()
        p.ExtraPoint()
        p.GameManagement()
        p.DisplayManagement()
        p.TGraphics()
        p.LogPlay(self.HTMDA, self.VTMDA, self.HTDMDA, self.VTDMDA)
        self._take_logs(p)
        GM['ConversionFlag'] = 0
        K['XPtFlag'] = 0
        GM['TDFlag'] = 0
        K['KickoffFlag'] = 1

    def go_for_two(self):
        GM, K = self.GM, self.Kicking
        p = self._play()
        p.GoForTwoPoints()
        p.GameManagement()
        p.DisplayManagement()
        GM['ConversionFlag'] = 0
        GM['TDFlag'] = 0
        K['KickoffFlag'] = 1

    def punt(self):
        GM, GO, K = self.GM, self.GameOptions, self.Kicking
        K['PlacementPunt'] = self.boxes["PlacementPunt"]
        GO['ForcePlay'] = self.boxes["ForcePlay"]
        if GM['Down'] < 4:
            bus.messages.append(("Button Pressed Incorrectly", "Punts Only Can Happen on 4th Downs"))
            return
        p = self._play()
        p.Punt()
        p.GameManagement()
        p.DisplayManagement()
        p.TGraphics()
        p.LogPlay(self.HTMDA, self.VTMDA, self.HTDMDA, self.VTDMDA)
        self._take_logs(p)
        K['PuntFlag'] = 0
        K['KickFlag'] = 0
        self.boxes["PlacementPunt"] = 0
        K['PlacementPunt'] = 0
        self._reset_boxes()

    def hto(self):
        self._play().TimeOutProcessing(1, 0)

    def vto(self):
        self._play().TimeOutProcessing(0, 1)

    def ot(self):
        if self.GM['OTFlag'] == 0:
            bus.messages.append(("Button Press Infraction", "Only press this button after the 1st OT Possession"))
            return
        p = self._play()
        p.OTManagement()
        p.DisplayManagement()

    # -------------------------------------------------------------- autoplay
    def game_over(self):
        return any("Game Over" in m[0] or "Game Over" in m[1] for m in self.messages)

    def auto_action(self):
        """A naive coach: kick when required, punt or try a FG on 4th and long."""
        gm = self.GM
        if gm['ConversionFlag'] == 1 or gm['TDFlag'] == 1:
            return "xpt"
        if self.Kicking['KickoffFlag'] == 1 or (gm['Quarter'] in (1, 3) and gm['TimeLeftinQuarter'] == 900):
            return "kickoff"
        if gm['Down'] == 4 and gm['YTG'] > 2:
            return "fg" if gm['YardLine'] >= 65 else "punt"
        if gm['OTFlag'] == 1 and gm['Down'] == 1 and gm['YardLine'] == 75 and gm['OTPossession'] > 1 \
                and not bus.board.get("play_call"):
            return "ot"
        return "call_play"

    def autoplay(self, max_actions=400):
        """Run the naive coach until the engine says the game is over."""
        n = 0
        while n < max_actions and not self.game_over():
            r = self.do(self.auto_action(), {})
            n += 1
            if r["error"]:
                return n, r
        return n, self.last_turn

    def quit(self):
        """QuitGame(): compile stats and write the log files.  Returns output file names."""
        bus.reset_turn()
        GM = self.GM
        if GM['HomeTeamScore'] > GM['VisitingTeamScore']:
            w, l, ws, ls = self.homeTeamName, self.visitingTeamName, GM['HomeTeamScore'], GM['VisitingTeamScore']
        else:
            w, l, ws, ls = self.visitingTeamName, self.homeTeamName, GM['VisitingTeamScore'], GM['HomeTeamScore']
        self.final_score = "%s: %s  %s: %s" % (w, ws, l, ls)
        error = None
        try:
            p = self._play()
            p.CompileStats(self.HTMDA, 0, self.HQBName, self.VQBName, self.HBackupQBName, self.VBackupQBName,
                           self.homeTeamName, self.visitingTeamName, self.final_score, self.VTDMDA)
            p.CompileStats(self.VTMDA, 1, self.HQBName, self.VQBName, self.HBackupQBName, self.VBackupQBName,
                           self.homeTeamName, self.visitingTeamName, self.final_score, self.HTDMDA)
            p.WriteLogFile(self.HTMDA, self.HomeTeamLogName)
            p.WriteLogFile(self.VTMDA, self.VisitingTeamLogName)
        except Exception:
            error = traceback.format_exc()
        self.finished = True
        outputs = (self.homeTeamName + "Stats.xlsx", self.homeTeamName + "DStats.xlsx",
                   self.visitingTeamName + "Stats.xlsx", self.visitingTeamName + "DStats.xlsx",
                   self.HomeTeamLogName, self.VisitingTeamLogName)
        # keep the outputs: move them into games/<id>/ so the next game does not overwrite them
        self.output_files = history.finish(self.id, GM['HomeTeamScore'], GM['VisitingTeamScore'], self.log, outputs)
        for scratch in (self.homeTeamName + "Log.xlsx", self.visitingTeamName + "Log.xlsx"):
            if os.path.exists(scratch):   # engine scratch logs, recreated by every start
                os.remove(scratch)
        self.save_journal("finished")
        return {"final_score": self.final_score, "files": self.output_files, "error": error, "id": self.id}

    def abandon(self):
        history.abandon(self.id, self.log)
        self.save_journal("abandoned")

    # ------------------------------------------------------------------ state
    def state(self):
        GM, b = self.GM, bus.board

        def clock():
            t = GM['TimeLeftinQuarter']
            return "%d:%02d" % (int(t) // 60, int(t) % 60)

        return {
            "home": self.homeTeamName, "visitor": self.visitingTeamName,
            "home_score": GM['HomeTeamScore'], "visitor_score": GM['VisitingTeamScore'],
            "quarter": GM['Quarter'], "clock": clock(),
            "down": GM['Down'], "ytg": b.get("ytg", GM['YTG']), "ballon": GM['AdjustedYardLine'],
            "home_timeouts": GM['HomeTimeouts'], "visitor_timeouts": GM['VisitorTimeouts'],
            "offense": "home" if GM['OffenseFlag'] == 0 else "visitor",
            "ot_series": b.get("ot_series", ""), "two_minute": b.get("two_minute", False),
            "play_call": b.get("play_call", ""), "last_result": b.get("last_result", ""),
            "ball_yardline": bus.ball["yardline"], "ball_offense_flag": bus.ball["offense_flag"],
            "boxes": dict(self.boxes), "log": list(self.log), "messages": list(self.messages),
            "finished": self.finished, "id": self.id, "can_undo": bool(self.journal),
        }
