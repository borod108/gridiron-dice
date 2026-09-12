#-----------------------------------------------------------------------------
# Name:        module1
# Purpose:
# Revison      4.0
# Author:      Rick
#
# Created:     10/02/2017
# Copyright:   (c) Rick 2017
# Licence:     <your license>
#-----------------------------------------------------------------------------


from Tkinter import *
#from LBwithWait_Window import * #May not be needed
from openpyxl import *
from ScoreboardClass import *

#import MDArray         #Not Needed.  Delete on Oct 28, 2021
import GeneralInfoFile #Reads a text file that contains general information such as the  
                                    #names of the last teams that played
import os                             #Allows to make MACoS system calls
import TestTurtleGraphics   #Draws football field
import Tkinter              #All of the widgets that we will need
#import Tkinter as tk

#Game Maintenance global objects
global down
global HomeTeamScore
global HomeTimeOuts             #Will be employed in revisions
global Quarter                  #1-4, OT
global OTSeries                 #An OT series is a possession by both teams during OT
global StartoftheGameFlag       #Used by Kickoff() for change of possession purpose
global TimeLeftinQuarter        #0 - 900 seconds.

#Log file global objects
global homeTeamPlayCount        #Used as an index into the home team play log
global PlayCount                #Aggregate for both teams.  Does not include kickoffs, FG, punts, Xpts
                                
#Scoreboard global objects
global OverTimeIndicator            #Indicates on the scoreboard that the game is in OT
global StartButtonHasBeenPressed    #Used to prevent incorrect button presses

#Team global objects
global homeTeam                     #IDs the home team worksheet
global homeTeamName                 #Working name of the home team
global TeamthatDoesNotHavetheBall   #Team on defense (IDs that team's worksheet)
global TeamWiththeBall              #Team on offense (IDs that team's worksheet)
global TwoPointFlag                 #Initialize but not currently used
global visitingTeam                 #IDs the visiting team worksheet
global visitingTeamName             #Working name of the visiting team
global visitingTeamPlayCount        #Used as an index into the visitors play log
global VisitingTeamScore            #Self expanatory
global VisitorTimeOuts              #Time out count in the current half for the visiting team
global YardLine                     #Where the ball is currently placed
global YardsToGo                    #Yards to go to a first down
#global x                           #Commented out and have not fully tested to make sure this variable is not used.
                                    #5/7/22

global GameHasStarted               #Flag to indicate the game has started and establishes the initial possession

GameHasStarted = 0                  #Initially, the game has not started

#Initialize variables

#Extra point variables
TwoPointFlag = 0                    #Only set during a 2-point try

#Game initialization variables
homeTeam = None                 #Set from Select Teams button which in turn is set from Team Selection file
#homeTeamName = "USC"           #Default if above method is not used
visitingTeam = None             #See above comments
#visitingTeamName = "UCLA"   

#Game management variables
down = 1                    #Always start on first down
OTSeries = 0                #Really should be set to 1 but the dictionary definition takes care of this
YardLine = 25               #Yard line will be next set by the kickoff
YardsToGo = 10              #Always start at 1st and 10

#Kickoff, FG and Punt variables
KickReturner = "X"              #Kick returner name.  Will be replaced by names from the worksheets
StartoftheGameFlag = 1          #Used to control possession
Startof3rdQuarterFlag = 0       #This flag is used to determine who receives the kickoff at the start of the 3rd qtr

#Scoreboard and logfile variables
homeTeamPlayCount = visitingTeamPlayCount = 0   #Used to compute tackle stats

HomeTimeOuts = 3                       #Each team has 3 timeouts per half
HomeTeamScore = VisitingTeamScore = 0  #Obviously, the game starts with the score 0-0
OverTimeIndicator = 1                  #Indicates when the game has gone into overtime
PlayCount = 0                          #Aggregate play countfor both teams.  Does not include kickoffs,
                                       #FG, punts, Xpts.  Unclear whether this variable is used at all.
                                       #Consider commenting out and seeing what happens
Quarter = 1                            #Game begins with quarter #1
TimeLeftinQuarter = 900                #In seconds, initialize to 15 minutes
VisitorTimeOuts = 3                    #Each team has 3 timeouts per half

#Define scoreboard objects
BallOn = Scoreboard()           #Displays the yardline where the ball is placed
BallOnLabel = Scoreboard()      #The label for this display indicator
Down = Scoreboard()             #Displays the down
DownLabel = Scoreboard()        #The label for this display indicator
HomeLabel = Scoreboard()        #The label for the indicator that displays the home team name
HomeOnOffense = Scoreboard()    #the indicator that is illuminated when the home team has the ball
HomeScore = Scoreboard()        #Displays the home team's current score
#ManualBox = Scoreboard()       #This was commented out on 5/20/22.  Note if error occurs, look here first
Quarter = Scoreboard()          #Identifies the current quarter
QuarterLabel = Scoreboard()     #The label for the displayed quarter
OTLabel = Scoreboard()          #The label for the display indicator that is illuminated when the game is in overtime
OverTime = Scoreboard()         #Display indicator that is illuminated when the game is in overtime
PlayCalledLabel = Scoreboard()  #Label for the field where the called play is displayed,  Not sure if this is used
PlayCalled = Scoreboard()       #Field that displays the current play that was called
PlayResult = Scoreboard()       #Field that displays the result of the play
PlayResult1 = Scoreboard()      #Not sure if this is used.  Investigate further
PlayResultLabel = Scoreboard()  #Label for the play result field
#TestLabel = Scoreboard()           #Good candidate to see if redundant-Remove 7/24/22
TimeLeftInQuarter = Scoreboard()    #Ranges from 0 - 15 minutes, format is minutes:seconds
TimoutsLabel = Scoreboard()         #Need one for home and visiting teams. Note spelling error
TimeoutsLeft = Scoreboard()         #This shows the # of tiimeouts left
TwoMinuteWarning = Scoreboard()     #Indicator on the scoreboard when there is < 2 minutes less
VisitorLabel = Scoreboard()         #Label indicating visitor side of the scoreboard
VisitorsOnOffense = Scoreboard()    #Indicator showing that the visiting team has the ball
VisitorsScore = Scoreboard()        #Obviously, the visiting team's score
YardsToGo = Scoreboard()            #Obviously, the yards to go
YardsToGoLabel = Scoreboard()       #The label for the obvious
StartButtonHasBeenPressed = 0       #0 = Game has not started, 1 = has started
    

#-----------------------------------------------------------------------------
# Function Name:    LoadTeams
# Purpose:          Based upon team name that is passed into this function load Excel spreadsheet for that team
# Inputs:           teamName: A string that is almost always an acronym for the team to keep things concise#
# Author:           Rick Burney
# Created:          12/13/2016
# Copyright:        (c) Rick 2016
#-----------------------------------------------------------------------------
def LoadTeams(teamName):
    
    import openpyxl #Using openpyxl Python library to handle all things EXCEL
        
    teamName += ".xlsx"                                     #All team workbooks have an xlsx extension
    wb = load_workbook(filename = teamName, data_only=True) #Load workbook
    ws = wb.active                                          #switch to active worksheet
    return ws                                               #Return the active ws to the calling procedure


#-----------------------------------------------------------------------------
# Function Name:    CoinToss
# Purpose:          See which team will receive the kickoff.  
# Inputs:           HomeTeam, VisitingTeam: These are pointers to the team worksheets
# Author:           Rick Burney
# Created:          12/13/2016
# Copyright:        (c) Rick 2016
#-----------------------------------------------------------------------------
def CoinToss(HomeTeam,VisitingTeam):
    
    import random   #Used for the random coin toss

    global TeamthatDoesNotHavetheBall   #Could have used the term defense
    global TeamwiththeBall              #Team that has the ball
    global TeamWiththeBallFlag          #0 if home team is on offense, 1 if visiting
    
    dice1 = random.randint(1,2)         #Flip a coin
    DisplayMsg = Scoreboard()           #Scoreboard object used to show who won the coin toss
    if dice1 == 1:                      #Home team receives the kickoff

        D = DisplayMsg.ResultDisplay(0,"Home Team will receive the kickoff",0,frame1)
        
        TeamWiththeBall = VisitingTeam          #This is counter-intuitive but will get corrected after the kickoff
        TeamWiththeBallFlag = 1                 #when there is a CoP.  Basically, the game starts w/ the kicking team
        TeamthatDoesNotHavetheBall = HomeTeam   # on offense
    else:                                       #
        D = DisplayMsg.ResultDisplay(0,
                            "Visiting Team will receive the kickoff",0,frame1)
        TeamWiththeBall = HomeTeam
        TeamWiththeBallFlag = 0
        TeamthatDoesNotHavetheBall = VisitingTeam
  
    
    return TeamWiththeBall


#-----------------------------------------------------------------------------
# Function Name:    Positions
# Purpose:          Parse the worksheet into positions and determine relevant
#                   indices
# Author:           Rick Burney
#
# Created:          12/13/2016
# Copyright:        (c) Rick 2016
#-----------------------------------------------------------------------------
def Positions(ws, TeamFlag):
    
    from openpyxl import load_workbook  #Library method that allows the loading of a file-specified workbook
    import PickAPlayer                  #Ultimately, will use this class to select players from the worksheet
    
    global QBsStartIndex        #This is where the QB stats are in the worksheet.  This is an EXCEL row
    global ReceiversStartIndex  #This is the row where receiver stats are in the stats worksheet
    global RunnersStartIndex    #This is the row where running back stats are in the stats worksheet
    
    PositionColumn = 1  #This is the EXCEL column where position headings are found.  Use to parse position groups
    QBNameColumn = 2    #This is the EXCEL column where the names of QBs are found

#Instantiate a class object that will identify the worksheet row where a particular position group or team stat is
#listed.
    RequestedRow = PickAPlayer.PickAPlayer(ws,0,0,0,0,0)    #Instantiate the class object
    RequestedRow.FindStats()                                #Go through the stat worksheets and identify the rows where
                                                            # the stats for position groups begin.  This is to
                                                            #eliminate the hard-coding of worksheet rows as constants.

    ReceiversStartRow = RequestedRow.ReceiversPosition + RequestedRow.Offset    #This is where the receiver stats are in
                                                                                #the worksheet
    QBsStartRow = RequestedRow.QBsPosition + RequestedRow.Offset                #This supercedes "QBStartIndex"
    RBsStartRow = RequestedRow.RunnersPosition + RequestedRow.Offset            #This supercedes "RunnersStartIndex"
    row_count = ws.max_row                                                      #Determine last row used in worksheet
    
#Parse through worksheet to identify rows where the position headings are located
    for i in range(1,row_count):
        if ws.cell(row=i, column = PositionColumn).value == 'Runners':      #RB
            RunnersRow = i
        if ws.cell(row=i, column = PositionColumn).value == 'QBs':          #QB
            QBsRow = i       
        if ws.cell(row=i, column = PositionColumn).value == 'Receivers':    #WR
            ReceiversRow = i
            
#Kickoffs, PATs, FGs and Punters
        if ws.cell(row=i, column=PositionColumn).value == 'Kickers':    #Punter's row are an offset from the kickers row
            KickersRow = i
        if ws.cell(row=i, column=PositionColumn).value == 'Dline':      #Not used.  Team defensive stats are used as are
            DlineRow = i                                                #individual sack stats
            if ws.cell(row=i, column=PositionColumn).value == 'LBs':    #Not used.  Team defensive stats are used as are
                LBsRow = i                                              #individual sack and interception stats
            if ws.cell(row=i, column=PositionColumn).value == 'DBs':    #Not used.  Team defensive stats are used as are
                DBsRow = i                                              #individual interception stats
                
#ID where the list of players who can carry the ball on a running play starts
    RunnersStartIndex = RBsStartRow 
    
#ID where the list of QBs start.  This includes the name of the QBs.
    QBsStartIndex = QBsStartRow
    QBName = ws.cell(row=QBsStartIndex, column = QBNameColumn).value
    BackupQBName = ws.cell(row=QBsStartIndex+1, column = QBNameColumn).value
    NumberOfReceivers = KickersRow - 2 - ReceiversRow                           #Use the categories above and below the
                                                                                #receivers group to determine the number
                                                                                #of receivers
    ReceiversStartIndex = ReceiversStartRow #Row where the receiver's stats begins
    QBNames = [QBName, BackupQBName]        #A list of the names of the starting and backup quarterback
    return QBNames


#-----------------------------------------------------------------------------
# Function Name:    StartGame
# Purpose:          Provides processing when the Start button is pressed.  This event will start the game
# Author:           Rick Burney
# Created:          11/09/2016
# Copyright:        (c) Rick 2016
#-----------------------------------------------------------------------------
def StartGame():
    
    import GameManager      #Contains the methods for calling plays, Run plays, pass plays
    import GeneralInfoFile  #Reads a text file that contains general information such as the names of the last teams
                            # that played
    import LogPlays         #Keeps a log of each play so that stats can be computed at the end of the game

   # import openpyxl         #May not be necessary.  Try out.  Commented out on July 7, 2022
    import PickAPlayer          #Method used to select a player for a particular stat such as a sack
    import STP                  #Make the class object instantiation from SelectTeams be global and this is where the
                                #teams that are selected are extracted
    import TestTurtleGraphics   #Draws football field
    #import Tkinter              #Used for widgets.  Commented out on July 16, 2022
    import tkMessageBox         #Used to send messages to the user
    import YardageTable         #Yardage tables that are used to determine the yardage on a play (in most cases)
    
#Game Maintenance
    #global down                #Commented out on July 19, 2022
    global OTFlag               #Indicates whether the game is in overtime or not
    #global Quarter             #Commented out on July 21, 2022
    #global Qtr                 #Commented out on July 22, 2022
    #global TimeLeftinQuarter   #Commented out on July 25, 2022
    #global YardsToGo           #Commented out on July 29, 2022
        
#Team Information and Stats
    global HDstats                      #home Team defensive stats
    #global homeTeam                    #Commented out on Aug 15, 2022
    global HomeTeamConferenceFactor     #Difference conferences are assigned a value that is used to augment or degrade
                                        #various offensive and defensive team stats
    #global homeTeamName                #Commented out on Aug 17, 2022
    global HomeTeamLogName              #File name for the home team log file
    #global x                           #Commented out on Sept 2, 2022
    
    #global TeamthatDoesNotHavetheBall  #Commented out Sept, 2022
    #global TeamWiththeBall             #Commented out on Sept 11, 2022
    global VDstats                      #Visiting team defensive stats
    #global visitingTeam                #Commented out on Sept 12, 2022
    global VisitingTeamConferenceFactor #Identifies the strength of conference for the visiting team
    #global visitingTeamName            #Commented out Oct 19, 2022
    global VisitingTeamLogName          #Name of the visiting team's log file for collecting statistics
    
#Player Information and Stats
    global HQBName                      #Name of the starting QB for the home team
    global VQBName                      #Name of the starting QB for the visiting team
    global HBackupQBName                #Name of the backup QB for the home team
    global VBackupQBName                #Name of the backup QB for the visiting team
    global HTeamStats                   #Currently, only used for home team penalty statistics (penalties/game)
    global VTeamStats                   #currently, only used for visiting team penalty statistics (penalties/game)
    
    global StartButtonHasBeenPressed #Used to prevent incorrect button presses
    
#Global Dictionaries that are used as data structures
    global GameOptions
    global GM
    global Kicking     #Dictionary for all things that use a foot
    
    global HTMDA                     #Offense, Home Team Memory Data Array
    global VTMDA                     #Offense, Visiting Team Memory Data Array
    global HTDMDA                    #Defense, Home Team Memory Data Array
    global VTDMDA                    #Defense, Visiting Team Memory Data Array
    global GameHasStarted            #Flag that shows that the Start Button has already been pressed
    
#Constants
    HomeTeamFlag = 0                #Constant that IDs the home team
    VisitingTeamFlag = 1            #Constant that IDs the visiting team

    
#Game Maintenance
    OTFlag = 0                      #Probably should comment out but won't know until we actually get into OT


#Player Information and Stats
    DIntPercColumn = 10             #The stats worksheet column where the interception percentage resides
    DIntPercRow = 46                #The stats worksheet row where the interception percentage risides
    DYPCTeamAveRow = 48

    
#Team Information and Stats
    ConferenceFactorColumn = 3      #The stats worksheet column where the conference factor resides
    ConferenceFactorRow = 90        #The initial stats worksheet row where the conference factor resides
    DRushTeamAveRow = 46            #The initial stats worksheet row where the team's yards/rush defense resides
    DRushTeamAveColumn = 3          #The stats worksheet column where the team's yards/rush defense resides
    DSacksPercRow = 47              #The initial stats worksheet row where the teams defensive sack percentage resides
    DSacksPerColumn = 6             #The stats worksheet column where the teams defensive sack percentage resides
    HTLeadingSpaces = ""            #Empty string that is used to create the home team name string
    PCTeamAveRow = 47               #The initial stats worksheet row where the teams defensive completion % resides
    PenaltyStatsColumn = 2          #The stats worksheet column where the team's penalties/game resides
    DFumRecoveryColumn = 10         #The stats worksheet column where the defense's fumble recovery probability resides
    DFumRecoveryRow = 47            #The stats worksheet row where the defense's fumble recovery probability resides
    TeamNameLength = 11             #Initial value.  This variable is ultimately used to display the team name on the
                                    #scoreboard
    TeamStatsRow = 90               #The stats worksheet row where, for now, the team's penalties/game resides
    #VTLeadingSpaces = ""           #Commented out on March 27, 2023
    StartButtonHasBeenPressed = 1    #Start Button has been pressed

#Need to write self.GM['TimeLeftinQuarter'] into TimeLeftinQuarter to make the next statement work.  Actually not sure
    #why this comment says what it says but this error check works so I will leave it alone
    if GameHasStarted == 1:
        tkMessageBox.showinfo("Error","Don't Press Start Button During Game")
        return


    
    visitingTeam = LoadTeams(visitingTeamName)  #Load the visiting team worksheet
    homeTeam = LoadTeams(homeTeamName)          #Load the home team worksheet

#Assign log file names.  During the game, the log data is stored in memory, at the end of the game, the log data is
#transferred to a file
    HomeTeamLogName = homeTeamName + "log.xlsx"                 #Home team log filename
    VisitingTeamLogName = visitingTeamName + "log.xlsx"         #Visiting team log filename
    homeTeamLeadingSpaces = TeamNameLength - len(homeTeamName)  #For display purposes, the team name that is displayed
    for i in range(1,homeTeamLeadingSpaces):                    #needs to be a fixed number of characters with leading
        HTLeadingSpaces += " "                                  #spaces
    HTLabel = HTLeadingSpaces + homeTeamName                    #For Scoreboard label for the home team
    HomeTeamLabel = VisitorTeamLabel = Scoreboard()             #Declare these labels as Scoreboard Class objects

 #Use WriteLabel method in Label Placement classes to write the team's label on the scoreboard
    HomeTeamLabelPlacement = HomeTeamLabel.WriteLabel(frame,HTLabel,10,2,"white","black",24,3,1,0,"SE")
    VisitorTeamLabelPlacement = VisitorTeamLabel.WriteLabel(frame,visitingTeamName,9,2,"white","black",24,3,6,0,"SW")

    x=LogPlays.ClearLog(homeTeamName,visitingTeamName)  #Call ClearLog from LogPlays class to clear both team's logs
    
#Initialize memory data arrays as empty lists.  Will create rows and append to lists after each play
    HTMDA = []     #Home team's offensive plays memory data array
    VTMDA = []     #Visiting team's offensive plays memory data array
    HTDMDA = []    #Home team's defensive plays memory data array
    VTDMDA = []    #Visiting team's defensive plays memory data array   

#Instantiate class objects that allows the rows on the worksheet where defensive team stats for both teams to be
#identified.  Note that the HTeam gets assigned to the visitor class object and vice versa to account for the fact that
#these are defensive stats so the opposite worksheet needs to be accessed.  Let's state this differently so when I look
#at this 20 years from now I will understand.  Say the home team is on offense.  ResultofthePlay() will use the
#offensive stats from the home team worksheet and Dstats which have to come from the visitor team worksheet. So VDstats
#is actually the home team defensive stats and vice versa.  Very confusing.

    RequestedHRow = PickAPlayer.PickAPlayer(homeTeam,0,0,0,0,0) #Instantiate a worksheet row object for the home team

#Go through the stat worksheets and identify the rows where the stats for position groups begin. This is to eliminate
#the hard-coding of worksheet rows as constants.  Do this for both teams
    RequestedHRow.FindStats()
    RequestedVRow = PickAPlayer.PickAPlayer(visitingTeam,0,0,0,0,0) 
    RequestedVRow.FindStats()
    VDRushTeamAveRow = RequestedVRow.DStatsPosition + RequestedVRow.Offset  #Defense against the run for both teams
    HDRushTeamAveRow = RequestedHRow.DStatsPosition + RequestedHRow.Offset  #Measured in yards/attempt

#Passes Completed Average against the defense.  This is the row on the worksheet that this stat is held
    VDPCTeamAveRow = RequestedVRow.DStatsPosition + RequestedVRow.DPCompPercOffset
    HDPCTeamAveRow = RequestedHRow.DStatsPosition + RequestedHRow.DPCompPercOffset
    VDYPCTeamAveRow = RequestedVRow.DStatsPosition + RequestedVRow.DPYACOffset      #Yards/catch pass defense for both
    HDYPCTeamAveRow = RequestedHRow.DStatsPosition + RequestedHRow.DPYACOffset      #teams
    
#Sacks by the defense are on the same row as completion % against the defense
    VDSacksPercRow = RequestedVRow.DStatsPosition + RequestedVRow.DPCompPercOffset
    HDSacksPercRow = RequestedHRow.DStatsPosition + RequestedHRow.DPCompPercOffset
    
#Ints are on the same row as yards per rush against the defense
    VDIntPercRow = RequestedVRow.DStatsPosition + RequestedVRow.Offset  
    HDIntPercRow = RequestedHRow.DStatsPosition + RequestedHRow.Offset  
    
#Fumble stats recovered by the defense are on the same row as completion % against the defense
    VDFumRecoveryRow = RequestedVRow.DStatsPosition + RequestedVRow.DPCompPercOffset
    HDFumRecoveryRow = RequestedHRow.DStatsPosition + RequestedHRow.DPCompPercOffset

#Stats come from individual team stat pages, usually with "team name" in the URL, but not always. Defensive team stats
#include rush defensive ave/rush, pass completion ave against the D, yards/completion against the D, Sack % for the D,
#% interceptions by the D, probability of a fumble by the D (on rushes, completions and sacks).  Extract defensive team
#stats for the visiting team
    VDstats = [homeTeam.cell(row = HDRushTeamAveRow,column = DRushTeamAveColumn).value,
               homeTeam.cell(row = HDPCTeamAveRow,column = DRushTeamAveColumn).value,
               homeTeam.cell(row = HDYPCTeamAveRow,column = DRushTeamAveColumn).value,
               homeTeam.cell(row = HDSacksPercRow,column = DSacksPerColumn).value,
               homeTeam.cell(row = HDIntPercRow,column = DIntPercColumn).value,
               homeTeam.cell(row = HDFumRecoveryRow,column = DFumRecoveryColumn).value]
                             
#Do the same for the home team defensive team stats
    HDstats = [visitingTeam.cell(row = VDRushTeamAveRow,column = DRushTeamAveColumn).value,
                visitingTeam.cell(row = VDPCTeamAveRow,column = DRushTeamAveColumn).value,
                visitingTeam.cell(row = VDYPCTeamAveRow,column = DRushTeamAveColumn).value,
                visitingTeam.cell(row = VDSacksPercRow,column = DSacksPerColumn).value,
                visitingTeam.cell(row = VDIntPercRow,column = DIntPercColumn).value,
                visitingTeam.cell(row = VDFumRecoveryRow,column = DFumRecoveryColumn).value]
    
#Identify row where penalties/game for both team's worksheets
    VDPenaltyStatsRow = RequestedVRow.TeamStatsPosition + RequestedVRow.Offset  
    HDPenaltyStatsRow = RequestedHRow.TeamStatsPosition + RequestedHRow.Offset  

#Extract penalties/game for both teams
    VTeamStats = [visitingTeam.cell(row = VDPenaltyStatsRow, column = PenaltyStatsColumn).value]
    HTeamStats = [homeTeam.cell(row = HDPenaltyStatsRow, column = PenaltyStatsColumn).value]
                                    
#For each team, determine the names of the starting and backup quarterbacks. Start by extracting a 2-field list
#containing the names of the starting and backup QB
    VQBNames = Positions(visitingTeam, VisitingTeamFlag)    #Extract QB name list for the visiting team
    VQBName = VQBNames[0]                                   #First name in the list is the starting QB
    VBackupQBName = VQBNames[1]                             #2nd name in the list is the backup QB
    HQBNames = Positions(homeTeam,HomeTeamFlag)             #Extract QB name list for the home team
    HQBName = HQBNames[0]                                   #1st name in the list is the starting QB
    HBackupQBName = HQBNames[1]                             #2nd name in the list is the backup QB
    
#Load conference factor for each team.  This allows for adjustment based upon the team's conference so that a team with
#great stats in a weak conference will be penalized.  The conference factors range from 4.5 to 5.  The conferences that
#have a conference factor of 5 are Big 10, PAC-12, ACC, SEC and Big-12. The conferences that have a conference factor
#b/n 4.5 and 5 are the AAC, Mountain West, MAC, Conference USA and Sun Belt.  The conference factor is found on the same
#row as the penalty stats, hence the row name
    HomeTeamConferenceFactor = homeTeam.cell(row = HDPenaltyStatsRow, column = ConferenceFactorColumn).value
    VisitingTeamConferenceFactor = visitingTeam.cell(row = VDPenaltyStatsRow, column = ConferenceFactorColumn).value

    TeamWiththeBall = CoinToss(homeTeam,visitingTeam)   #Do the coin toss
    down = 1                                            #First and 10
    YardsToGo = 10
    Quarter = 1
    T = Scoreboard()                                                            #Declare a scoreboard object
    Ds = T.ResultDisplay(3,"15:00",0,frame)                                     #Initialize timekeeping on display
    
#Here is where we stored the teams playing to the general info file for 
#retrieval for the next time the simulator is run
    GeneralInfoFile.WriteGeneralInfo(homeTeamName,visitingTeamName)   

#Dictionary definition.  GM stands for Game Management and is a dictionary that is global and is used to pass important
#game info throughout the app.  Some of the dictionary keys are obvious but some are not.  TeamWiththeBall is the
#the offense, TeamthatDoesNotHavetheBall is the defense. OffenseFlag deals with kickoffs which prior to the kick,
#there is no team with the ball yet.  OffenseFlag defines who kicks off and who receives.
    GM = {'Down': 1, 'YTG': 10, 'YardLine' : 25, 'AdjustedYardLine' : 25, 'TimeLeftinQuarter' : 900, 'Quarter' : 1,
          'Offense' : TeamWiththeBall, 'Defense' : TeamthatDoesNotHavetheBall, 'OffenseFlag' : TeamWiththeBallFlag,
          'OldYardLine': 25, 'OldAdjustedYardline' : 25, 'OldDown' : 1, 'OldYTG' : 10, 'OldTimeLeftInQuarter' : 900,
          'OldQuarter' : 1, 'HomeTeamScore' : 0, 'VisitingTeamScore' : 0, 'CoPFlag' : 0, 'OldHTS' : 0, 'OldVTS' : 0,
          'Touchback' : 0,'ConversionOnDowns' : 0, 'TimeCode' : 0, 'OTFlag' : 0, 'OTCounter' : 0, 'SafetyFlag' : 0,
          'DTDFlag': 0, 'homeTeamPlayCount' : 0, 'visitingTeamPlayCount' : 0, 'OTSeries' : 1, 'OTPossession' : 1,
          'TDFlag': 0, 'UntimedDownFlag' : 0, 'HomeTimeouts' : 3, 'VisitorTimeouts' : 3, 'OldTimeLeftInQuarter' : 0,
          'TurnoverTD' : 0,'ClockRunning' : 0, 'ConversionFlag' : 0}

#GameOptions is a dictionary that is global and is for strategy-based approaches to offense and defense.  Teams that are
#run-oriented can be directed to favor runs, for example.  A team can put in 2nd string players if they are winning
#"big."  The QB can take a knee. 'ForcePlay' is used for debugging with extra code inserted into other methods
    GameOptions = {'RunCentric': 0, 'PassCentric' : 0, 'Hup' : 0, 'SpiketheBall' : 0, 'HailMary' : 0, 'Blowout' : 0,                      
                   'ForcePenalty' : 0, 'QBTakesaKnee' : 0, 'ForcePlay' : 0}

#Kicking is a dictionary that contains global objects for used by all methods associated with kicking or punting.  For
#example, Kickflag is set when a kicking play is about to be executed and is reset when a non-kicking play is to be
#executed. StartoftheGame is set to detect an accidental press of the Start Button.  Startofthe3rdQuarter identifies
#when the 3rd quarter starts to determine who kicks off to whom (and sets the ThridQuarterKickoffTeam and
#ThirdQuarterTeamReceivingtheKick objects). OnsideKick is set when an onside kick is to be executed. SquibKick is set
#when a squib kick is to be executed.  Free kick is set immediately after a safety. ThirdQuarterOffenseFlag determines
#who is on offense at the start of the 3rd quarter. XPtFlag is set when an extra point is being
#executed. KickoffFlag is set when the kick is a kickoff. FGFlag is set when the kick is a field goal. PlacementPunt is
#set when a placement punt is requested. PuntFlag is set when the kick is a punt.  TwoPointConvFlag is set when a 2-
#point conversion is attempted. ThirdQuarterKickoffRow indicates the row of the kicker that kicks off at the start of
#the 3rd quarter and is a STARTING POINT TO FIND A BUG REGARDING KICKING
    Kicking = {'KickFlag' : 0, 'StartoftheGame' : 0, 'Startofthe3rdQuarter' : 0, 'OnsideKick' : 0, 'SquibKick' : 0,
               'FreeKick' : 0, 'ThirdQuarterKickoffTeam' : TeamthatDoesNotHavetheBall,
               'ThirdQuarterTeamReceivingtheKick' : TeamWiththeBall, 'ThirdQuarterOffenseFlag' : 0, 'XPtFlag' : 0,
               'KickoffFlag' : 0, 'FGFlag' : 0, 'PlacementPunt' : 0, 'PuntFlag' : 0, 'TwoPointConvFlag' : 0,
               'ThirdQuarterKickoffRow' : 0}
    Kicking['StartoftheGame'] = StartoftheGameFlag  #Notifies the Kicking Method that this is the opening kickoff
    GameHasStarted = 1                              #Flag that is used to ignore presses of the Start Button after the
                                                    #the game has started
    
#-----------------------------------------------------------------------------
# Function Name:    SelectTeams
# Purpose:          Called when the Select Teams button is pressed.  Will open an Excel spreadsheet that allows the user
#                   a much friendlier GUI then that of the Rev 1.0 version.  Would have preferred to use a class object
#                   to do this but Python appears to have a bug and the widgets flash after the class has been destroyed
#                   but a TK messagebox is invoked
# Author:           Rick Burney
# Created:          7/3/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def SelectTeams():

    import openpyxl         #Uses an EXCEL team selection worksheet to select teams
    import os               #Operating system specific library

    global homeTeamName     #These two objects will hold the name of the home 
    global visitingTeamName #and visitor teams - strings - they are the acronyms of the team names
    ACCSearchColumn = 3     #ACC column for the team selection worksheet
    AmericanSearchColumn=39 #American Conference column for the team select ws
    Big10SearchColumn = 7   #Big-10 Conference column for the team select. ws
    Big12SearchColumn = 11  #Big-12 Conference column for the team select. ws
    CUSASearchColumn = 43   #Conference USA column for the team select. ws
    PAC12SearchColumn = 15  #PAC-12 Conferencecolumn for the team select. ws
    IndySeachColumn = 19    #Independents column for the team selection ws
    MACSearchColumn = 23    #Mid-America Conference column for the ws
    SECSearchColumn = 27    #SEC column for the team selection worksheet
    
    SunBeltSearchColumn = 31        #Sun Belt Conference column for the ws
    MountainWestSearchColumn = 35   #Mountain West Conf column for the ws
    FCSSearchColumn = 47                   #FCS column for the ws
    
#Organize all the columns into a search list
    SearchColumns = [ACCSearchColumn,Big10SearchColumn,Big12SearchColumn, PAC12SearchColumn,IndySeachColumn,
                     MACSearchColumn, SECSearchColumn,SunBeltSearchColumn, MountainWestSearchColumn,
                     AmericanSearchColumn, CUSASearchColumn, FCSSearchColumn]

#The team selection worksheet has a fixed name in the current directory. All team workbooks have an .xlsx extension.
    SelectTeamWorksheet = "Team Selection.xlsx"
    wb = load_workbook(filename = SelectTeamWorksheet, data_only=True)  #Load Team Selection Workbook
    ws = wb.active                                                      #switch to active worksheet
    for i in range(2,ws.max_row+1): #Search thru the team selection worksheet to determine the home and visitor teams
        for j in SearchColumns:
            
            if str(ws.cell(row=i, column=j).value).upper() == "H": #The home team is marked with an "H"
                homeTeamName = ws.cell(row=i, column=j-2).value

            if str(ws.cell(row=i, column=j).value).upper() == "V": #The visitor team is marked with a "V"
                visitingTeamName = ws.cell(row=i, column=j-2).value
    

#-----------------------------------------------------------------------------
# Function Name:    CallPlay1
# Purpose:          Provides processing when the Call Play button is pressed and is redesigned from CallPlay for
#                   Revision 3 to support extensive use of classes
# Author:           Rick Burney
# Created:          7/4/2017
# Copyright:        (c) Rick 2017
#-------------------------------------------------------------------------------
def CallPlay1():
    
    import GameManager          #GameManager is a collection of functions that manage the execution of play calls
    import PickAPlayer          #Class used to find stat locations in worksheet
    import Play                 #Class where the play is executed
    import TestTurtleGraphics   #Draw the ball position on the football field
    import Tkinter              #Tkinter is a Python library that can be used to construct basic graphical user
                                # interface (GUI) applications.
    import tkMessageBox         #The tkMessageBox module is used to display message boxes in applications.
    
    
#Dictionaries
    global GameOptions #Data structure that holds the checkbutton options
    global GM          #The main data structure used to keep track of the game
    global Kicking     #Dictionary for all things that use a foot
    global PlayerStats #Data structure with pointers to individual player stats
    global TeamStats   #Data structure with pointers to team stats
    
#Player Information and Stats
    global BallCarrier          #The name of the ball carrier.  It could be a receiver.
    global BallCarrierAve       #Yards per carry or catch
    global HDstats              #See StartGame for detailed definitions for Dstats
    global QBsStartIndex        #Used to find where the QBs are on the ws
    global RunnersStartIndex    #Used to find where the RBs are on the ws
    global ReceiversStartIndex  #Same for the receiver
    global VDstats              #See StartGame for detailed definitions for Dstats
    global homeTeam             #ID's home team stats worksheet
    global visitingTeam         #ID's visiting team stats worksheet
    
    global StartButtonHasBeenPressed #Used to prevent incorrect button presses
    
#Memory data arrays are used for realtime data logging
    global HTMDA                     #Offense, Home Team Memory Data Array
    global VTMDA                     #Offense, Visiting Team Memory Data Array
    global HTDMDA                    #Defense, Home Team Memory Data Array
    global VTDMDA                    #Defense, Visiting Team Memory Data Array
    

#Check for incorrect button presses.  Can't call a play before the game has started and you can't call a play at the
#beginning of the 1st or 3rd quarter before you have done a kickoff. Can't call a play after a touchdown if an extra
#point or a "go for 2" has not been attempted
    if StartButtonHasBeenPressed == 0:
        tkMessageBox.showinfo("Button Press Infraction","Game must be started before you can press this button")
        return
    if ((GM['Quarter'] == 1) or (GM['Quarter'] == 3)) and \
       (GM['TimeLeftinQuarter'] == 900):
        tkMessageBox.showinfo("Button Press Infraction", "Press the Kickoff Button")        
        return
    
    if GM['ConversionFlag'] == 1:   #TD just scored, must do a conversion before proceeding
        tkMessageBox.showinfo("Button Press Infraction",
        "Press Extra Point, Go For 2 button or Kickoff button as appropriate")
        return
    if Kicking['KickoffFlag'] == 1:
        tkMessageBox.showinfo("Button Press Infraction", "Only press the Kickoff Button")
        return
    GameOptions['Blowout'] = 0  #Assume no blowout, user needs to keep blowout checkbutton checked
    CoPFlag = 0                 #Always want this flag to be 0 until there is a legitimate change of possession

#Create a dictionary of pointers to certain stats.  I believe this has been overcome by PickAPlayer
    PlayerStats = {'RunnerSI': RunnersStartIndex,  'ReceiverSI' : ReceiversStartIndex,  'QBSI' : QBsStartIndex}

#Read state of checkbuttons    
    GameOptions['RunCentric'] = RunOriented.get()     #Increases probability of a running play
    GameOptions['PassCentric'] = PassOriented.get()   #Increases probability of a pass play
    GameOptions['Hup'] = HurryUp.get()                #Hurry up offense
    GameOptions['SpiketheBall'] = Spike.get()         #QB spikes the ball
    GameOptions['HailMary'] = HailMaryPass.get()      #Hail Mary pass
    GameOptions['ForcePlay'] = ForcePlay.get()        #Diagnostic
 
 #Instantiate a Play class object for debugging   
    DebugMethod = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,frame,Kicking)
    
    HomeTeamBlowout = HTBlowout.get()                   #Home team is blowing out visitors
    VisitingTeamBlowout = VTBlowout.get()               #Visiting team is blowing out the home team
    GameOptions['QBTakesaKnee'] = QBTakesaKnee.get()    #QB takes a knee checkbox
    
#This next section checks the RunOriented and the PassOriented checkbuttons and if both are set, this is an error and
#both will be reset
    if (GameOptions['RunCentric'] == 1) and (GameOptions['PassCentric'] == 1):
        PassOriented.set(0)
        RunOriented.set(0)

#If either or both of the Blowout Checkbuttons are checked, use the backup players as appropriate
    if (GM['OffenseFlag'] == 0) and (HomeTeamBlowout == 1):
        GameOptions['Blowout'] = HomeTeamBlowout
    if (GM['OffenseFlag'] == 1) and (VisitingTeamBlowout == 1):
        GameOptions['Blowout'] = VisitingTeamBlowout

#The following dictionary object holds the team stats used by this program for both the offense and the defense
    TeamStats = {'HomeDStats': HDstats, 'VisitorDStats' : VDstats, 
                 'HomePenaltiesPerGame' : HTeamStats[0], 
                 'VisitorPenaltiesPerGame' : VTeamStats[0], 
                 'HTCF' : HomeTeamConferenceFactor, 
                 'VTCF' : VisitingTeamConferenceFactor, 
                 'homeTeamName' : homeTeamName, 
                 'visitingTeamName' : visitingTeamName, 
                 'PreCoPOffenseFlag' : GM['OffenseFlag'], 
                 'HomeTeamBlowout' : HomeTeamBlowout, 
                 'VisitingTeamBlowout' : VisitingTeamBlowout, 
                 'homeTeam' : homeTeam, 'visitingTeam' : visitingTeam}

    if GameOptions['Blowout'] == 1:
        PlayerStats['QBSI'] = QBsStartIndex+ 1  #Use the 2nd string QB stats in the case of a blowout
    else:
        PlayerStats['QBSI'] = QBsStartIndex     #Otherwise use the 1st string QB stats

#******************************DEBUG*******************************************************
    GameOptions['ForcePenalty'] = GameOptions['ForcePlay'] #Use code like this as a template for debuggin
    #print(GameOptions['ForcePlay'])
    #print(GameOptions['ForcePenalty'])
#******************************DEBUG*******************************************************

#Call a play by instantiating a play object
    PlayExecution = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,frame,Kicking)
    PlayExecution.OnOffenseIndication()
    PlayExecution.PlayCall()                            #Call a play
    Play = PlayExecution.Play                           #Assign the play call to an object
    PenaltyList = PlayExecution.PreSnapPenaltyTest()    #Test for a penalty including pre-snap                                                        
    PenaltyList = PlayExecution.PenaltyList             #If penalty, set this flag, otherwise, set to false
    if PenaltyList[1] == 0:                             #No pre-snap Penalty        
        PlayExecution.PlayResult()                      #Run the play
        ResultList = PlayExecution.ResultList           #Extract the partial play results

#Can't have a pass interference on a sack        
        if ((PenaltyList[0] == "Pass Interference on the Defense") or \
            (PenaltyList[0] == "Pass Interference on the Offense")) and (ResultList[2] == "Sack"):
            PenaltyList[5] = 0

#As long as there is not a turnover, update the yardline once the play has been completely resolved.  Turnovers are
#handled differently
        if (PlayExecution.FumbleResult[0] == 0) and (PlayExecution.IntResult[5] == 0):
            PlayExecution.UpdateYardLine()      
        if PenaltyList[5] == 1:                             
            PlayExecution.PenaltyProcessing()   #Now check for penalties. PenaltyList is set by the play result
        PlayExecution.GameManagement()          #Update Game Management and store in global dictionary
        GM = PlayExecution.GM                   #If the play results in a change of possession, update the CoP flag in
        CoPFlag = GM['CoPFlag']                 #the GM dictionary

        if (CoPFlag == 1) and (GM['OTFlag'] == 0):  #If the CoP flag is set during regulation execute a CoP and reset
            PlayExecution.CoP()                    #all checkboxes
            ResetAllCheckBoxes()
    
    PlayExecution.DisplayManagement()  #Update the Scoreboard accordingly
    
    if PenaltyList[0] == "Personal Foul on the Defense, Hands to the Face":    #Lie in wait for this #condition to occur
        
        tkMessageBox.showinfo("Debug","Debug condition 0 - check debug window")        
        PlayExecution.DebugAfterThePlay(1)   
    #if (GM['SafetyFlag'] == 1) and (PenaltyList[5] == 1):
        #tkMessageBox.showinfo("Debug","Debug condition 1 - check debug window")        
        #PlayExecution.DebugAfterThePlay(1)   
        
#Log play that just was run
    PlayExecution.LogPlay(HTMDA,VTMDA,HTDMDA,VTDMDA)
    HTMDA = PlayExecution.HTMDA                         #Update global copies
    VTMDA = PlayExecution.VTMDA                         #of logs
    HTDMDA = PlayExecution.HTDMDA
    VTDMDA = PlayExecution.VTDMDA
    GM['SafetyFlag'] = 0            #Reset flags
    GM['ScoreFlag'] = 0 
    PlayExecution.TGraphics()
    return


#-------------------------------------------------------------------------------
# Function Name:    HTimeout
# Purpose:          Timeout processing when the HTO button is pressed.  
# Author:           Rick Burney
# Created:          11/26/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def HTimeout():

#Dictionaries
    global GameOptions #Data structure that holds the checkbutton options
    global GM          #The main data structure used to keep track of the game
    global Kicking     #Dictionary for all things that use a foot
    import Play        #Class where the play is executed
    global PlayerStats #Data structure with pointers to individual player 
                           #stats
    global TeamStats   #Data structure with pointers to team stats
    
    HomeTimeOut = 1     #Timeout called by the home team
    VisitorTimeOut = 0
    TO = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,frame,Kicking)
    
    TO.TimeOutProcessing(HomeTimeOut,VisitorTimeOut)    #Call a timeout
    
     
#-----------------------------------------------------------------------------
# Function Name:    VTimeout
# Purpose:          Timeout processing when the VTO button is pressed.  
# Author:           Rick Burney
# Created:          11/26/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def VTimeout():

#Dictionaries
    global GameOptions #Data structure that holds the checkbutton options
    global GM          #The main data structure used to keep track of the game
    global Kicking     #Dictionary for all things that use a foot
    import Play                 #Class where the play is executed
    global PlayerStats #Data structure with pointers to individual player 
                           #stats
    global TeamStats   #Data structure with pointers to team stats
        
    HomeTimeOut = 0     #Because there are 2 different buttons for timeouts, 
    VisitorTimeOut = 1  #there needs to be two command methods, one for the
                        #home team, the other for the visiting team.  These 2
                        #flags are set such that this IDs the timeout as 
                        #having come from a press of the visiting team timeout
                        #button

#Instantiate a timeout object        
    TO = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,frame,Kicking)
    
    TO.TimeOutProcessing(HomeTimeOut,VisitorTimeOut)    #Call a visiting team 
                                                        #timeout    
   

#-----------------------------------------------------------------------------
# Function Name:    ResetAllCheckBoxes
# Purpose:          Resets all checkbodes
# Author:           Rick Burney
# Created:          3/20/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def ResetAllCheckBoxes():
    
    import Tkinter            #Needed to access checkbox
    import tkMessageBox #Needed to send message to the screen alerting the user about 
                                       #something
        
    Squib.set(0)               #Reset kickoff checkboxes.  Squib kickoff
    Onside.set(0)             #Onside kick
    FK.set(0)                    #FK is the free kick checkbox, used after a safety
    PlacementPunt.set(0) #Reset punt checkboxes
    RunOriented.set(0)     #Reset play call checkboxes.  Makes offense run-oriented
    PassOriented.set(0)    #Makes offense pass-oriented
    HurryUp.set(0)           #Hurry up offense
    Spike.set(0)                #QB spikes the ball
    HailMaryPass.set(0)    #Hali Mary Pass


#-----------------------------------------------------------------------------
# Function Name:    OTSetup1 - this seems to be broken 5/18/21
# Purpose:          Sets up and runs multiple series of OT.  OT follows the following sequence.  
#                        The first OT consists of Series 1 and 2, with 1 team on offense on Series 1 
#                        and the other team on offense on Series 2.  For Series 3, the team on 
#                        offense on Series 2 is on offense and then the team on offense on Series 1 
#                   will be on offense for Series 4.  This
#                   alternating pattern will continue but after every even-
#                   numbered series, the score will be evaluated and if the
#                   score is no longer tied, the team with more points wins 
#                   and the game is over.  Each team starts from the opposing
#                   team's 25.  After the 2nd OT (Series 1-4), all TDs will be
#                   followed by a 2-point conversion, no Xtra points.  Only 
#                   globals will be passed to and from this function.
# Author:           Rick Burney
# Created:          8/10/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def OTSetup1():
    import GameManager          #Not sure if this is needed
    import Play                 #Class where the play is executed
    import TestTurtleGraphics   #Draw the ball position on the football field
    import Tkinter              #Message boxes are used to communicate OT
    import tkMessageBox         #status with the user
    
#Dictionaries
    global GameOptions #Data structure that holds the checkbutton options
    global GM          #The main data structure used to keep track of the game
    global Kicking     #Dictionary for all things that use a foot
    global PlayerStats #Data structure with pointers to individual player 
                       #stats
    global TeamStats   #Data structure with pointers to team stats
    
#The program automatically enters OT.  After the first possession, the user
#progresses through the OT series by pressing the OT Button.  After the 1st
#series, OTFlag will be set
    if GM['OTFlag'] == 0:
        tkMessageBox.showinfo("Button Press Infraction","Only press this \
        button after the 1st OT Possession")
        return
    
#Create an OT object
    TheOT = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,
                              frame,Kicking)
    TheOT.OTManagement()                        #Perform OT management
    TheOT.DisplayManagement()                   #Display OT status
    

#-----------------------------------------------------------------------------
# Function Name:    Kickoff1
# Purpose:          New Kickoff method that uses the Play Class
# Author:           Rick Burney
# Created:          7/22/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def Kickoff1():

    import Play     #Class where the play is executed
    import Tkinter
    import tkMessageBox    
    
#Dictionaries
    global GameOptions #Data structure that holds the checkbutton options
    global GM          #The main data structure used to keep track of the game
    global Kicking     #Dictionary for all things that use a foot
    global PlayerStats #Data structure with pointers to individual player 
                       #stats
    global TeamStats   #Data structure with pointers to team stats
    
    global StartoftheGameFlag
    
#Player Information and Stats
    global BallCarrier          
    global BallCarrierAve
    global HDstats              #See StartGame for detailed definitions for
                                #Dstats
    global QBsStartIndex        #Used to find where the QBs are on the ws
    global RunnersStartIndex    #Used to find where the RBs are on the ws
    global ReceiversStartIndex  #Same for the receiver
    global VDstats
    
#Error checking
    global StartButtonHasBeenPressed
    
    
#Memory data arrays are used for realtime data logging
    global HTMDA   #Offense, Home Team Memory Data Array, Kickoffs go here
    global VTMDA   #Offense, Visiting Team Memory Data Array, Kickoffs go here
    global HTDMDA  #Defense, Home Team Memory Data Array, 
    global VTDMDA  #Defense, Visiting Team Memory Data Array
    
    global HTeamStats
    global VTeamStats

    GM['ConversionFlag'] = 0                          #Reset extra point flag

#Extract penalty stats for both teams. 
    TeamStats = {'HomePenaltiesPerGame' :  HTeamStats[0], 
                 'VisitorPenaltiesPerGame' : VTeamStats[0]}
    

    GameOptions['ForcePlay'] = ForcePlay.get()        #Diagnostic
    
    PlayerStats = None  #Have not been defined yet so this will tide us over
#    TeamStats = None
    ThirdQuarterTeamWiththeBallFlag = None

#Press kickoff button only after the game has been started
    if StartButtonHasBeenPressed == 0:
        tkMessageBox.showinfo("Button Press Infraction","Game must be started\
        before you can press this button")
        return
    Kicking['KickoffFlag'] = 1  #IDs this play as a kickoff.  Have subsequently set from Xpt and 
                                            #GoFor2 to prevent incorrect key presses
    
#Create a kickoff (play) instance
    TheKickoff = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,frame,Kicking)
    
    Kicking['SquibKick'] = Squib.get()          #Read kickoff checkbuttons  
    Kicking['OnsideKick'] = Onside.get()
    Kicking['FreeKick'] = FK.get()

#Stupid little trick to see if more than 1 CheckBox is checked
    NumKickCBsChecked = Kicking['SquibKick'] + Kicking['OnsideKick'] + \
        Kicking['FreeKick']   
    
    if NumKickCBsChecked > 1: #If more than 1 CB is checked, reset all and
        Squib.set(0)          #notify the user
        Onside.set(0)                       
        FK.set(0)
        Kicking['SquibKick'] = Kicking['OnsideKick'] = Kicking['FreeKick'] = 0

        tkMessageBox.showinfo("Dummy","Only one of the kicking checkboxes \
        can be checked at once so all were reset and ignored")
        
   
    K_O = TheKickoff.Kickoff()  #Kick off
    TheKickoff.LogPlay(HTMDA,VTMDA,HTDMDA,VTDMDA)
    HTMDA = TheKickoff.HTMDA                         #Update global copies
    VTMDA = TheKickoff.VTMDA                         #of logs
    HTDMDA = TheKickoff.HTDMDA
    VTDMDA = TheKickoff.VTDMDA
    TheKickoff.GameManagement()     #Update everything based upon the kickoff
    TheKickoff.DisplayManagement()  #results and then display 
    TheKickoff.TGraphics()          #Update the field display
    GM['ConversionFlag'] = 0        #This is only necessary after a free kick
                                    #that is after a safety
    
    ResetAllCheckBoxes()
    
        


#-----------------------------------------------------------------------------
# Function Name:    FG1
# Purpose:          Receives the yardline from the calling program and 
#                   determines if the kick is good or not
# Author:           Rick Burney
# Created:          7/31/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def FG1():
 
    import Play         #The primary class used for all play procesing
    import tkMessageBox #Used to convey error messages and debugging info
    
    global GameOptions #Data structure that holds the checkbutton options
    global GM          #The main data structure used to keep track of the game
    global Kicking     #Dictionary for all things that use a foot
    global PlayerStats #Data structure with pointers to individual player 
                       #stats
    global TeamStats   #Data structure with pointers to team stats
    
#Memory data arrays are used for realtime data logging
    global HTMDA   #Offense, Home Team Memory Data Array, Kickoffs go here
    global VTMDA   #Offense, Visiting Team Memory Data Array, Kickoffs go here
    global HTDMDA  #Defense, Home Team Memory Data Array, 
    global VTDMDA  #Defense, Visiting Team Memory Data Array
    

    GameOptions['ForcePlay'] = ForcePlay.get()        #Diagnostic

    if GM['ConversionFlag'] == 1:
        #tkMessageBox("Error","Conversion Flag is set and should not be")
        tkMessageBox.showinfo("Button Press Infraction",
        "Press Extra Point, Go For 2 button or Kickoff button as appropriate")
        return
 

 #Create an Field Goal (Play) instance
    TheFG = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,
                              frame,Kicking)
    
    F_G = TheFG.FieldGoal()     #Kick the field goal, update the scoreboard
    if Kicking['FGFlag'] == 1:
        
        TheFG.GameManagement()
        
        TheFG.DisplayManagement()  #and display the results
        TheFG.TGraphics()          #Update the field display
        
        TheFG.LogPlay(HTMDA,VTMDA,HTDMDA,VTDMDA)    #Log the result of the FG
        HTMDA = TheFG.HTMDA                         #Update global copies
        VTMDA = TheFG.VTMDA                         #of logs
        HTDMDA = TheFG.HTDMDA
        VTDMDA = TheFG.VTDMDA
        Kicking['FGFlag'] = 0     #Reset field goal flag

    Kicking['KickFlag'] = 0 #Reset this flag, field goal has been executed



#-----------------------------------------------------------------------------
# Function Name:    XPt1
# Purpose:          New Extra Point method that uses the Play Class
# Author:           Rick Burney
# Created:          7/27/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def XPt1():
    import Play         #Main class used for plays
    import tkMessageBox #Used for error message to the user, usually due to an
                        #incorrect button press    
    global GameOptions #Data structure that holds the checkbutton options
    global GM          #The main data structure used to keep track of the game
    global Kicking     #Dictionary for all things that use a foot
    global PlayerStats #Data structure with pointers to individual player 
                       #stats
    global TeamStats   #Data structure with pointers to team stats

#Memory data arrays are used for realtime data logging
    global HTMDA   #Offense, Home Team Memory Data Array, Kickoffs go here
    global VTMDA   #Offense, Visiting Team Memory Data Array, Kickoffs go here
    global HTDMDA  #Defense, Home Team Memory Data Array, 
    global VTDMDA  #Defense, Visiting Team Memory Data Array
    
    global StartButtonHasBeenPressed    #Used to tell the user to start the 
                                        #game before pressing any buttons

#Prevents the user from inadvertantly pressing the extra point button.  
#Additional safeguards have been implemented in the Play Class
    if StartButtonHasBeenPressed == 0:
        tkMessageBox.showinfo("Button Press Infraction","Game must be started\
        before you can press this button")
        return

#If in overtime, on and after the 3rd series, teams must go for 2
    if (GM['OTFlag'] == 1) and (GM['OTSeries'] >= 3):
        tkMessageBox.showinfo("Overtime","On or after the 3rd series, and after a TD, the scoring team must go for 2 points")
        return
    GameOptions['ForcePlay'] = ForcePlay.get()        #Diagnostic
      
#Create an Extra Point (Play) instance
    TheXPt = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,frame,Kicking)
    
    X_Pt = TheXPt.ExtraPoint()       #Kick the extra point, update the scoreboard
    TheXPt.GameManagement()     #and display the results by creating a play class object   
    TheXPt.DisplayManagement()   
    TheXPt.TGraphics()                  #Update the field display
    
    TheXPt.LogPlay(HTMDA,VTMDA,HTDMDA,VTDMDA)    #Log the result of the XPt
    HTMDA = TheXPt.HTMDA                                           #Update global copies of logs
    VTMDA = TheXPt.VTMDA                         
    HTDMDA = TheXPt.HTDMDA
    VTDMDA = TheXPt.VTDMDA
    GM['ConversionFlag'] = 0      #Reset conversion flag
    Kicking['XPtFlag'] = 0            #Reset extra point flag    
    GM['TDFlag'] = 0                   #Now you can reset the TD flag.  This is done to ensure that
                                                 #the only time the Extra Point button is pressed is
                                                 #immediately after a TD
    Kicking['KickoffFlag'] = 1  #Set to prevent CallPlay from being pressed when a kickoff is
                                             #expected
    
    
#-----------------------------------------------------------------------------
# Function Name:    GoForTwo1
# Purpose:     Determines if a 2-Point conversion is successful.  Right now, for reasons of
#                   convenience, an actual play will not be run.  Also, can't find team stats so the
#                   the average success rate in college football is 42 %.  Recoded for the Play Class
#                   rate in college football is 42 %.   
#                   No player or play will be identified
# Author:      Rick Burney
# Created:     8/2/2017
# Copyright: (c) Rick 2017
#-----------------------------------------------------------------------------
def GoForTwo1():
    
    import Play                 #Now uses Play Class
    import tkMessageBox #Not used right now but who knows   
    global GameOptions   #Data structure that holds the checkbutton options
    global GM                   #The main data structure used to keep track of the game
    global Kicking             #Dictionary for all things that use a foot
    global PlayerStats        #Data structure with pointers to individual player stats
    global TeamStats         #Data structure with pointers to team stats
    
    TwoPtConv = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,
                              frame,Kicking)
    
    TPC = TwoPtConv.GoForTwoPoints()    #Try for a 2 point conversion
    TwoPtConv.GameManagement()          #Update the display
    TwoPtConv.DisplayManagement()
    GM['ConversionFlag'] = 0        #Reset conversion flag
    GM['TDFlag'] = 0                    #Reset here for the same reason as the
                                        #extra point
    Kicking['KickoffFlag'] = 1  #Set to prevent CallPlay from being pressed
                                #when a kickoff is expected
 

#-----------------------------------------------------------------------------
# Function Name:    Punt1
# Purpose:          Executes a punt using the Play Class
# Author:           Rick Burney
# Created:          8/3/2017
# Copyright:        (c) Rick 2017
#-----------------------------------------------------------------------------
def Punt1():
 
    import Play         #Now uses the Play Class
    import tkMessageBox #Needed so that the user can be informed that punts
                        #only occur on 4th downs
    
    global GameOptions #Data structure that holds the checkbutton options
    global GM          #The main data structure used to keep track of the game
    
#Memory data arrays are used for realtime data logging
    global HTMDA   #Offense, Home Team Memory Data Array, Kickoffs go here
    global VTMDA   #Offense, Visiting Team Memory Data Array, Kickoffs go here
    global HTDMDA  #Defense, Home Team Memory Data Array, 
    global VTDMDA  #Defense, Visiting Team Memory Data Array
 
    global Kicking     #Dictionary for all things that use a foot
    global PlayerStats #Data structure with pointers to individual player 
                       #stats
    global TeamStats   #Data structure with pointers to team stats
    
    Kicking['PlacementPunt'] = PlacementPunt.get()  #Read check buttons
    GameOptions['ForcePlay'] = ForcePlay.get()      #Debug check button

#Only punt on 4th downs    
    if GM['Down'] < 4:
        tkMessageBox.showinfo("Button Pressed Incorrectly",
                              "Punts Only Can Happen on 4th Downs")
        return
    
#Instantiate a punt object    
    ThePunt = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,
                              frame,Kicking)
    
    P = ThePunt.Punt()          #Punt the ball, Play class handles blocks, no
    ThePunt.GameManagement()    #returns, fair catches and punt returns
    ThePunt.DisplayManagement() #Display the result
    ThePunt.TGraphics()         #Update the field display
    
    ThePunt.LogPlay(HTMDA,VTMDA,HTDMDA,VTDMDA)
    HTMDA = ThePunt.HTMDA                         #Update global copies
    VTMDA = ThePunt.VTMDA                         #of logs
    HTDMDA = ThePunt.HTDMDA
    VTDMDA = ThePunt.VTDMDA
    Kicking['PuntFlag'] = 0     #Reset kick flags
    Kicking['KickFlag'] = 0
    PlacementPunt.set(0)                            #Reset placement punt 
    Kicking['PlacementPunt'] = PlacementPunt.get()  #check button
    ResetAllCheckBoxes()                            #Reset all other
                                                    #checkboxes as appropriate

    

#-----------------------------------------------------------------------------
# Function Name:    QuitGame
# Purpose:          Exits the simulator
# Author:           Rick Burney
#
# Created:          11/10/2016
# Copyright:        (c) Rick 2016
#-----------------------------------------------------------------------------
def QuitGame():
   
    import Play     #Class where the play is executed 
    import turtle   #Need this to close turtle window

    import Tkinter            #Needed for message boxes.  
    import tkMessageBox #A Message box is used to display the final score

#Team Information and Stats
    global homeTeamName       #Used for the final score Message Box that is displayed at the
    global HomeTeamLogName #end of the game
    global visitingTeamName     
    global VisitingTeamLogName

#Dictionaries
    global GameOptions #Data structure that holds the checkbutton options
    global GM                 #The main data structure used to keep track of the game
    global Kicking          #Dictionary for all things that use a foot
    global PlayerStats     #Data structure with pointers to individual player stats
    global TeamStats      #Data structure with pointers to team stats
    global HTMDA          #Home and visiting team logs that are used to compile
    global VTMDA          #the final stats
    global HTDMDA       #Home and visiting team logs that are used to compile
    global VTDMDA       #the stats such as funbles and interceptions (from the defense's
                                    #perspective)
  
#Declare a Play class object that will be used to compile stats for both teams
    Log = Play.Play(GM,PlayerStats,GameOptions,TeamStats,frame1,frame,Kicking)
    
#Compile offensive and defensive stats for both teams. Two arguments that are included in 
#the call are required because if the last play of the game is a kick, then these two dictionary 
#items are not defined because they are only defined in CallPlay1 for a very good reason 
#that I can't remember.  So it is the global objects that are passed.
    
#Display a message box with the final score
    if GM['HomeTeamScore'] > GM['VisitingTeamScore']:
        WinningTeam = homeTeamName
        LosingTeam = visitingTeamName
        WinningScore = str(GM['HomeTeamScore'])
        LosingScore = str(GM['VisitingTeamScore'])
    else:
        WinningTeam = visitingTeamName
        LosingTeam = homeTeamName
        WinningScore = str(GM['VisitingTeamScore'])
        LosingScore = str(GM['HomeTeamScore'])
    
    FinalScore = WinningTeam + ": " + WinningScore + "  " + \
        LosingTeam + ": " + LosingScore 
    tkMessageBox.showinfo("Final Score",FinalScore)
    Log.CompileStats(HTMDA, 0,HQBName,VQBName,HBackupQBName,VBackupQBName, 
                     homeTeamName, visitingTeamName,FinalScore,VTDMDA)
    Log.CompileStats(VTMDA, 1,HQBName,VQBName,HBackupQBName,VBackupQBName, 
                     homeTeamName, visitingTeamName,FinalScore,HTDMDA)
    Log.WriteLogFile(HTMDA,HomeTeamLogName)
    Log.WriteLogFile(VTMDA,VisitingTeamLogName)
    

    root.destroy()  #Quit the application 
    turtle.bye()    #Close the turtle window

#***************************MAIN LOOP****************************************

#Program starts here
root = Tkinter.Tk(  )

#Have the display grab the focus
os.system('''/usr/bin/osascript -e 'tell app "Finder" to set frontmost of \
process "Python" to true' ''')

#Draw scoreboard
root.geometry("1350x600")
root.configure(background='black')
frame = Frame(root, height = 500, width = 1350, bg = "black")
frame1 = Frame(root, height = 175, width = 1350, bg = "black")
frame.pack()
frame1.pack()

#Define checkbox variables and initialize global variables
var = IntVar()
varm = IntVar()
autorun = IntVar()
autorun = 0
manual = IntVar()
manual = 0




#Place the buttons on the Scoreboard and ID the command method that will 
#process the button
#Start Button
ButtonName = Tkinter.Button(frame, text="Start",command=StartGame).grid(row=1,
                            column=2,padx=20,pady=20)

#Kickoff Button
ButtonName = Tkinter.Button(frame, text="Kickoff",command=Kickoff1).\
                                          grid(row=1,column=3,padx=20,pady=20)

#Play Execution Buttons.  "command" IDs the routine that processes the Field Goal button
ButtonName = Tkinter.Button(frame, text="FG",command=FG1).grid(row=1,column=4,
                                                               padx=20,pady=20)

#Extra Point Button
ButtonName = Tkinter.Button(frame, text="XPt",command=XPt1).grid(row=1,column=5,
                                                                 padx=20,pady=20)

#Go For 2 button
ButtonName = Tkinter.Button(frame, text="Go For 2",command=GoForTwo1).\
                                          grid(row=1,column=6,padx=20,pady=20)

#Punt button
ButtonName = Tkinter.Button(frame, text="Punt",command=Punt1).grid(row=1,column=7,
                                                                   padx=20,pady=20)

#Call Play button.  This button is used the most
ButtonName = Tkinter.Button(frame, text="Call Play",
                            command=CallPlay1).grid(row=1,column=9,padx=20,pady=20)

#Home Team Timeout Button
ButtonName = Tkinter.Button(frame, text="HTO",
                            command=HTimeout).grid(row=1,column=8,padx=20,pady=20)

#Visiting Team Timeout Button
ButtonName = Tkinter.Button(frame, text="VTO",
                            command=VTimeout).grid(row=1,column=10,padx=20,pady=20)

#Quit Button.  This wraps everything up and invokes the compilation of stats
ButtonName = Tkinter.Button(frame, text="Quit",command=QuitGame).grid(row=6,
                            column=7,padx=20,pady=20)

#Overtime button.  This inititiates the OT process and sets up the OT operating conditions
ButtonName = Tkinter.Button(frame, text="OT",command=OTSetup1).grid(row=5,
                                                                    column=10,padx=0,pady=0)

#Loads in the Team Selection worksheet where the user has already selected the home and 
#visiting teams
ButtonName = Tkinter.Button(frame, text="SelectTeams",
                            command=SelectTeams).grid(row=1,column=1,padx=0,pady=0)
                        
#This label is placed below the Scoreboard indicator that IDs the home team                   
HomeLabelPlacement = HomeLabel.WriteLabel(frame,"          HOME",10,2,"white","black",24,3,
                                          1,0,"SE")

#This is the current home team score placed near the home team ID on the Scoreboard 
HomeScorePlacement = HomeScore.WriteLabel(frame,"00",3,1,"black","white",36,4,1,2,"NE")

#This is a green indicator that is illuminated green when the home team has the ball
HomeOnOffensePlacement = HomeOnOffense.WriteLabel(frame,"",2,1,"green","green",8,4,2,0,
                                                  "W")

#This displays the time left in the quarter
TLIQPlacement = TimeLeftInQuarter.WriteLabel(frame,"00:00",6,2,"white","black",36,3,4,0,"")

#Displays the label over which the visiting team's shortened name is displayed
VisitorLabelPlacement = VisitorLabel.WriteLabel(frame,"VISITORS",9,2,"white","black",24,3,6,0,
                                                "SW")

#Displays the label over which the home team's remaining timeouts for the half are displayed
TimoutsLabelPlacement = TimoutsLabel.WriteLabel(frame,"TIMEOUTS",9,2,"white","black",20,
                                                5,1,0,"W")

#Displays the label over which the visiting team's remaining timeouts for the half are 
#displayed
TimoutsLabelPlacement = TimoutsLabel.WriteLabel(frame,"TIMEOUTS",9,2,"white","black",20,
                                                5,7,0,"W")

#Displays the label over which the quarter is displayed
QtrLabelPlacement = QuarterLabel.WriteLabel(frame,"Qtr",4,2,"white","black",24,4,3,0,"E")

#Displays the label over which the the "Visitors are on Offense" indicator is displayed
VisitorsOnOffensePlacement = VisitorsOnOffense.WriteLabel(frame,"",2,1,"black","grey",8,4,
                                                          6,0,"W")

#Displays the label over which the visitor's score is displayed
VisitorsScorePlacement = VisitorsScore.WriteLabel(frame,"00",3,1,"black","white",36,4,6,5,"")

#Displays the quarter
QuarterPlacement = Quarter.WriteLabel(frame,"1",3,1,"black","white",24,4,4,0,"")

#Place Timeout boxes for both the home team and the visitors
TimeoutsPlacement = TimeoutsLeft.WriteLabel(frame,"3",3,1,"white","black",20,5,2,0,"W")
TimeoutsPlacement = TimeoutsLeft.WriteLabel(frame,"3",3,1,"white","black",20,5,8,0,"W")

#Displays the label over which the down is displayed
DownLabelPlacement = DownLabel.WriteLabel(frame,"DOWN",5,1,"white","black",22,7,1,0,"")

#Displays the down
DownLocation = Down.WriteLabel(frame,"1",2,2,"white","black",24,7,1,0,"E")

#Displays the Yards to Go
YardsToGoLocation = YardsToGo.WriteLabel(frame,"10",3,2,"white","black",24,7,4,0,"W")

#Displays the label over which the yards to go are displayed
YardsToGoLabelLocation = YardsToGoLabel.WriteLabel(frame,"To Go",6,2,"white","black",24,7,
                                                   3,2,"E")

#Displays the yardline for the current line of scrimmage
BallOnLocation = BallOn.WriteLabel(frame,"25",3,2,"white","black",24,7,6,0,"W")

#Displays the label over which the yardline for the current line of scrimmage is displayed
BallOnLabelLocation = BallOnLabel.WriteLabel(frame,"Ball is On",9,2,"white","black",24,7,5,0,
                                             "W")

#Indicates that the game is in overtime
OTL = OTLabel.WriteLabel(frame,"OverTime",9,2,"white","black",24,6,8,0,"")
PlayResultLabelLocation = PlayResultLabel.WriteLabel(frame1,"Play Result",11,
                                                     2,"white","black",24,4,3,
                                                        0,"E")
PlayResultLocation = PlayResult.WriteLabel(frame1,"",  
                                           85,2,"black","white",24,1,4,5,"W")
PlayResult1Location = PlayResult1.WriteLabel(frame1,
                                             "",85,2,
                                             "black","white",24,4,4,5,"W")

RunOriented = IntVar()
TimeIsShort = Tkinter.Checkbutton(frame, text="Run-Oriented",
                                  variable = RunOriented ).grid(row=2,
                                            column=9, padx=0,pady=0,
                                            sticky = "")
PassOriented = IntVar()
TimeIsShort = Tkinter.Checkbutton(frame, text="Pass-Oriented",
                                  variable = PassOriented ).grid(row=2,
                                            column=10, padx=0,pady=0,
                                            sticky = "")
HurryUp = IntVar()
TimeIsShort = Tkinter.Checkbutton(frame, text="Hurry-Up",
                                  variable = HurryUp ).grid(row=3,
                                            column=9, padx=0,pady=0,
                                            sticky = "")
Spike = IntVar()
SpikeCheckBox = Tkinter.Checkbutton(frame, text="SpiketheBall",
                                  variable = Spike).grid(row=3,
                                            column=10, padx=0,pady=0,
                                            sticky = "")
Squib = IntVar()
SquibCheckBox = Tkinter.Checkbutton(frame, text="Squib Kick",
                                  variable = Squib).grid(row=2,
                                            column=2, padx=0,pady=0,
                                            sticky = "")
Onside = IntVar()
OnsideKickCheckBox = Tkinter.Checkbutton(frame, text="Onside Kick",
                                  variable = Onside).grid(row=2,
                                            column=3, padx=0,pady=0,
                                            sticky = "")
PlacementPunt = IntVar()
PlacementPuntCheckBox = Tkinter.Checkbutton(frame, text="Placement Punt",
                                  variable = PlacementPunt).grid(row=2,
                                            column=7, padx=0,pady=0,
                                            sticky = "")
FK = IntVar()
FreeKickCheckBox = Tkinter.Checkbutton(frame, text="Free Kick",
                                  variable = FK).grid(row=2,
                                            column=4, padx=0,pady=0,
                                            sticky = "")

TMW = Scoreboard()
TMWIndicatorLabel = TMW.WriteLabel(frame,"2 Min Warn",10,1,"white",
                                                "black",12,7,8,0,"W")

HailMaryPass = IntVar()
HailMaryPassCB = Tkinter.Checkbutton(frame, text="Hail Mary Pass",
                                  variable = HailMaryPass).grid(row=4,
                                            column=10, padx=0,pady=0,
                                            sticky = "")
HTBlowout = IntVar()
Blowout = Tkinter.Checkbutton(frame, text="HT Blowout",
                                  variable = HTBlowout).grid(row=6,
                                            column=1, padx=0,pady=0,
                                            sticky = "")
VTBlowout = IntVar()
Blowout = Tkinter.Checkbutton(frame, text="VT Blowout",
                                  variable = VTBlowout).grid(row=6,
                                            column=10, padx=0,pady=0,
                                            sticky = "")
ForcePlay= IntVar()
FP = Tkinter.Checkbutton(frame, text="Force",
                                  variable = ForcePlay).grid(row=6,
                                            column=2, padx=0,pady=0,
                                            sticky = "")


QBTakesaKnee = IntVar()
QBTakesaKneeCB = Tkinter.Checkbutton(frame, text="Take a Knee",
                                  variable = QBTakesaKnee).grid(row=4,
                                            column=9, padx=0,pady=0,
                                            sticky = "")


TeamList = GeneralInfoFile.ReadGeneralInfo()    #Read the last 2 teams that 
                                                #played
homeTeamName = TeamList[0]
visitingTeamName = TeamList[1]
DF = TestTurtleGraphics.DrawField()             #Draw the football field


if __name__ == "__main__":
    root.mainloop(  )