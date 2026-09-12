def LoadTeamMethod():
    

#from openpyxl import *
    from openpyxl import load_workbook
    import random
    wb = load_workbook(filename = 'USC.xlsx', data_only=True)  #Load workbook

    ws = wb.active                              #switch to active worksheet            

    #Extract delimiters (indices of start of list)
    #Find EOF
    row_count = ws.max_row

    for i in range(1,row_count):
        if ws.cell(row=i, column = 1).value == 'Runners':
            RunnersRow = i
        #print RunnersRow
        if ws.cell(row=i, column = 1).value == 'QBs':
            QBsRow = i
        #print QBsRow
        if ws.cell(row=i, column = 1).value == 'Receivers':
            ReceiversRow = i
        #print WRsRow
        if ws.cell(row=i, column=1).value == 'Kickers':
            KickersRow = i
        if ws.cell(row=i, column=1).value == 'Dline':
            DlineRow = i
        if ws.cell(row=i, column=1).value == 'LBs':
            LBsRow = i
        if ws.cell(row=i, column=1).value == 'DBs':
            DBsRow = i

    #Develop ranges and indices
    #Runners Range
    RunnersRange = range(RunnersRow+1,QBsRow-1)
    NumberOfRunners = QBsRow - 2 - RunnersRow
    RunnersStartIndex = RunnersRange[0]

    #QBs Range
    QBsRange = range(QBsRow+1,ReceiversRow-1)
    NumberOfQBs = ReceiversRow - 2  - QBsRow
    QBsStartIndex = QBsRange[0]

    #Receivers Range
    ReceiversRange = range(ReceiversRow+1,KickersRow-1)
    NumberOfReceivers = KickersRow - 2 - ReceiversRow
    ReceiversStartIndex = ReceiversRange[0]

    print(ws["H2"].value)
    return None




    
